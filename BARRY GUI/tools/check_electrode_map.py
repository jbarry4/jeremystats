"""check_electrode_map.py -- probes.py still says what the workbook says.

`backend/probes.py` holds the DEWEY channel map as Python literals, the same
way it holds the H10-D's geometry, so that module reads no files. That is the
right call -- a probe template is not data that changes under you -- but it
means the transcription can drift from `electrode_map.xlsx` and nothing would
notice until a matrix came out with two regions swapped.

So this reads the workbook and diffs it against the literals.

It also re-states the disagreement with the older map every run, so nobody has
to rediscover it: the 64-channel HOF map put left ACC on 21-22 and left OFC on
23-24, and this workbook has them the other way round. The workbook is the
truth (confirmed 2026-09-24) -- but a connectivity figure made from the J1/J2
pilot has those two regions swapped relative to anything made here, and that
is worth saying out loud rather than filing away.

    python tools\\check_electrode_map.py
"""
from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, APP)

from backend import probes  # noqa: E402

BOOK = os.path.join(APP, "electrode_map.xlsx")
SHEET = "Channel Map"


def from_workbook():
    """{csc number: (sheet label, full name, hemisphere)} as the sheet has it."""
    try:
        import openpyxl
    except ImportError:
        print("openpyxl is not installed, so the workbook cannot be read.")
        print("  pip install openpyxl")
        return None
    if not os.path.exists(BOOK):
        print("workbook not found: %s" % BOOK)
        return None

    wb = openpyxl.load_workbook(BOOK, data_only=True)
    if SHEET not in wb.sheetnames:
        print("workbook has no sheet %r (has %s)" % (SHEET, wb.sheetnames))
        return None

    out = {}
    for row in wb[SHEET].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            ch = int(row[0])
        except (TypeError, ValueError):
            continue            # the note row at the bottom
        out[ch] = (str(row[1] or "").strip(),
                   str(row[2] or "").strip(),
                   str(row[3] or "").strip())
    return out


def from_code():
    """The same shape, built from the literals in probes.py.

    `sheet_name` is what the workbook says; `name` is what the app shows.
    They differ where a label was normalised so the twelve regions read the
    same way everywhere, and it is the WORKBOOK's spelling that is compared
    here -- checking the display name instead would just be the check
    marking its own homework.
    """
    out = {}
    for reg in probes.DEWEY_REGIONS:
        for ch in reg["csc"]:
            out[int(ch)] = (reg["sheet_label"],
                            reg.get("sheet_name") or reg["name"],
                            reg["hemisphere"])
    return out


def main():
    sheet = from_workbook()
    if sheet is None:
        return 2
    code = from_code()

    problems = []
    for ch in sorted(set(sheet) | set(code)):
        want = sheet.get(ch)
        got = code.get(ch)
        if want is None:
            problems.append((ch, "probes.py claims it; the workbook does not",
                             None, got))
        elif got is None:
            problems.append((ch, "the workbook has it; probes.py does not",
                             want, None))
        elif want[0].lower() != got[0].lower():
            problems.append((ch, "region label", want, got))
        elif want[1].lower() != got[1].lower():
            problems.append((ch, "full name", want, got))
        elif want[2].lower() != got[2].lower():
            problems.append((ch, "hemisphere", want, got))

    print("workbook : %s" % BOOK)
    print("channels : %d in the sheet, %d in probes.py"
          % (len(sheet), len(code)))
    print("regions  : %d" % len(probes.DEWEY_REGIONS))
    print("disagreements : %d" % len(problems))

    renamed = [r for r in probes.DEWEY_REGIONS
               if r.get("sheet_name") and r["sheet_name"] != r["name"]]
    if renamed:
        print("\nLabels the app shows differently from the workbook:")
        for r in renamed:
            print("  CSC %-7s %r -> %r"
                  % (",".join(str(c) for c in r["csc"]),
                     r["sheet_name"], r["name"]))

    flagged = [r for r in probes.DEWEY_REGIONS if r.get("was")]
    if flagged:
        print("\nRegions the older 64-channel map disagreed about:")
        for r in flagged:
            print("  %-10s CSC %s" % (r["region"],
                                      ", ".join(str(c) for c in r["csc"])))
            print("     %s" % r["was"])

    if problems:
        print("\nFAIL -- probes.py and the workbook disagree:\n")
        for ch, what, want, got in problems:
            print("  CSC %-3d %s" % (ch, what))
            print("     workbook  %s" % (want,))
            print("     probes.py %s" % (got,))
        return 1

    print("\nOK -- probes.py matches the workbook, channel for channel.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
