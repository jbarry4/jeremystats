# -*- coding: utf-8 -*-
"""Import what the lab's spreadsheets know into BARRY.

Three facts live in two Excel files and nowhere else, which means they are
true only for whoever has the file open:

    bad channels    Mouse Recording Sessions.xlsx, one row per recording
    layer per channel   the feeder sheet's `location` column
    hemisphere      the feeder sheet's `side` column

Run with no arguments it changes nothing and prints what it would do. That
is the default on purpose: this writes to real recordings, and the first
version of any importer is wrong in a way you only see in the diff.

    python tools/import_feeder.py                 # say what would change
    python tools/import_feeder.py --apply         # do it
    python tools/import_feeder.py --apply --only bad,side,layers

The join between the sheets and BARRY is (mouse, session), which is the only
thing both sides carry -- BARRY keys on a gid derived from the folder, and
the sheets have never heard of it. Where BARRY holds more than one recording
for a (mouse, session), the row is reported and skipped rather than guessed
at: writing a layer sheet onto the wrong recording is worse than not writing
it at all.
"""
import argparse
import collections
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                          # noqa: E402

from backend import layers as layersmod                  # noqa: E402

BAD_XLSX = r"D:\PTEN\Mouse Recording Sessions.xlsx"
FEED_XLSX = r"D:\PTEN\EEGData_FeederSheet_PTEN_with_DKO.xlsx"

OUT = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")


def say(msg=""):
    OUT.write(msg + "\n")
    OUT.flush()


def clean(v):
    return "" if v is None else str(v).strip()


def sess_key(v):
    """'m11s10' -> (11, 10)."""
    m = re.match(r"^m(\d+)s(\d+)", clean(v).lower().replace(" ", ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def chan_of(fname):
    """'CSC12.ncs' -> 12."""
    m = re.search(r"(\d+)", clean(fname))
    return int(m.group(1)) if m else None


# --------------------------------------------------------------------------
# The layer vocabulary
#
# Every value in the sheet's `location` column already has a home in
# layers.REGIONS except two, and both are real rather than typos:
#
#   THAL   twelve deep channels of m2s3. Thalamus -- below the hippocampus
#          and outside it, but emphatically not "out of brain", which is the
#          only other place it could have gone.
#   DG2    m30's lower-blade dentate with no sublayer given. The vocabulary
#          has three lower-blade layers and no way to say "lower blade,
#          unspecified", which is exactly what this row means.
#
# Mapping either onto something that already exists would be recording a
# guess as though it were data.
# --------------------------------------------------------------------------
EXTRA_REGIONS = [
    {"id": "thal", "name": "THAL", "color": "#b48ead", "note": "thalamus"},
    {"id": "dg2", "name": "DG2", "color": "#7fb069",
     "note": "dentate, lower blade, unspecified"},
]

LOCATION_TO_REGION = {
    "CA1": "ca1_sr",          # the vocabulary's ca1_sr is named "CA1"
    "CA1 SR": "ca1_sr",
    "CA1 SP": "ca1_sp",
    "CA1 SO": "ca1_so",
    "CA1 SLM": "ca1_slm",
    "DG OML1": "dg_oml1",
    "DG MML1": "dg_mml1",
    "DG GCL1": "dg_gcl1",
    "HIL": "hil",
    "DG GCL2": "dg_gcl2",
    "DG MML2": "dg_mml2",
    "DG OML2": "dg_oml2",
    "DG": "dg",
    "THAL": "thal",
    "DG2": "dg2",
}


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------
def read_bad():
    """(mouse, session) -> sorted bad channel numbers."""
    wb = openpyxl.load_workbook(BAD_XLSX, read_only=True, data_only=True)
    got = {}
    for row in wb["Sheet1"].iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * 5
        m = re.match(r"^m?(\d+)$", clean(row[0]).lower())
        if not m or not clean(row[1]):
            continue
        key = (int(m.group(1)), int(float(clean(row[1]))))
        got[key] = sorted({int(x) for x in re.findall(r"\d+", clean(row[4]))})
    wb.close()
    return got


def read_feeder():
    """(mouse, session) -> {"side": 'L'/'R', "chans": {chan: location}}.

    Every tab is read and the answers are compared rather than the last one
    winning. The eight tabs overlap heavily -- three of them describe the
    same first batch -- and an import that quietly takes whichever sheet it
    happened to read last is an import nobody can check. As it happens they
    agree everywhere; the check stays because that is not a promise the
    spreadsheet has made.
    """
    wb = openpyxl.load_workbook(FEED_XLSX, read_only=True, data_only=True)
    cells = collections.defaultdict(dict)      # (key, chan) -> {sheet: (side, loc)}
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 3:
            continue
        head = [clean(c).lower() for c in rows[1]]

        def col(*names):
            for n in names:
                if n in head:
                    return head.index(n)
            return None

        i_sess, i_file = col("sess"), col("file")
        i_side, i_loc = col("side"), col("location")
        if i_sess is None or i_file is None or i_loc is None:
            continue
        for r in rows[2:]:
            r = list(r) + [None] * 24
            key, ch = sess_key(r[i_sess]), chan_of(r[i_file])
            if key is None or ch is None:
                continue
            cells[(key, ch)][name] = (
                clean(r[i_side]).upper() if i_side is not None else "",
                clean(r[i_loc]))
    wb.close()

    conflicts = []
    out = collections.defaultdict(lambda: {"side": "", "chans": {}})
    for (key, ch), bysheet in cells.items():
        sides = {v[0] for v in bysheet.values() if v[0]}
        locs = {v[1] for v in bysheet.values() if v[1]}
        if len(sides) > 1 or len(locs) > 1:
            conflicts.append((key, ch, dict(bysheet)))
            continue
        if locs:
            out[key]["chans"][ch] = locs.pop()
        if sides:
            out[key]["side"] = sides.pop()
    return dict(out), conflicts


def barry_index(store):
    """(mouse, session) -> [session records]."""
    by = collections.defaultdict(list)
    for s in store.all_sessions():
        mouse, sess = s.get("mouse"), s.get("session")
        if mouse is None or sess is None:
            continue
        try:
            by[(int(mouse), int(sess))].append(s)
        except (TypeError, ValueError):
            continue
    return by


# --------------------------------------------------------------------------
# The three jobs
# --------------------------------------------------------------------------
def do_bad(store, index, sheet, apply_it, mode="report"):
    """Bad channels, with a deliberate choice about removals.

    BARRY already holds bad channels for these recordings, from the Toothy
    workbook. Against that, this spreadsheet adds nothing at all: every
    difference is a channel BARRY calls bad and the sheet does not. So
    "confirm bad channels from here" cannot be done by writing the column
    over the top -- the entire effect would be to un-flag thirteen channels
    across eleven recordings.

    Un-flagging is not a neutral edit. A channel marked bad is left out of
    what people look at and of what they compute; clearing the mark puts
    whatever was wrong with it back into the analysis, quietly, in recordings
    somebody may already have drawn conclusions from.

    So the default reports and writes nothing. `union` adds without ever
    removing, which is the safe direction. `replace` treats the sheet as
    authoritative, and is a decision somebody makes on purpose.
    """
    say("\n" + "=" * 70)
    say("BAD CHANNELS   (mode: %s)" % mode)
    say("=" * 70)
    same = changed = added = skipped = 0
    for key in sorted(sheet):
        want = sheet[key]
        recs = index.get(key) or []
        if len(recs) != 1:
            say("  m%-3d s%-3d  SKIPPED: BARRY has %d recordings for this"
                % (key[0], key[1], len(recs)))
            skipped += 1
            continue
        rec = recs[0]
        have = sorted(int(b) for b in (rec.get("bad_channels") or []))
        if have == want:
            same += 1
            continue
        final = sorted(set(have) | set(want)) if mode == "union" else want
        drops = sorted(set(have) - set(final))
        gains = sorted(set(final) - set(have))
        if final == have:
            same += 1
            if set(want) != set(have):
                say("  m%-3d s%-3d  %-22s   sheet says %s   [kept as is]"
                    % (key[0], key[1], have, want))
            continue
        say("  m%-3d s%-3d  %-22s -> %-18s%s%s"
            % (key[0], key[1], have or "(none)", final,
               ("  +%s" % gains) if gains else "",
               ("  DROPS %s" % drops) if drops else ""))
        if have:
            changed += 1
        else:
            added += 1
        if apply_it and mode != "report":
            store.set_bad_channels(
                {"gid": rec.get("gid"), "mouse": key[0], "session": key[1],
                 "key": rec.get("key"), "loose_key": rec.get("loose_key"),
                 "label": rec.get("label")},
                final,
                note="From Mouse Recording Sessions.xlsx")
    say("\n  unchanged %d   newly set %d   replaced %d   skipped %d"
        % (same, added, changed, skipped))
    return {"same": same, "added": added, "changed": changed,
            "skipped": skipped}


def do_side(store, index, feed, apply_it):
    say("\n" + "=" * 70)
    say("HEMISPHERE")
    say("=" * 70)
    same = added = changed = skipped = 0
    for key in sorted(feed):
        want = feed[key]["side"]
        if want not in ("L", "R"):
            continue
        recs = index.get(key) or []
        if len(recs) != 1:
            say("  m%-3d s%-3d  SKIPPED: %d recordings" % (key[0], key[1], len(recs)))
            skipped += 1
            continue
        rec = recs[0]
        have = clean(rec.get("hemisphere")).upper()
        if have == want:
            same += 1
            continue
        say("  m%-3d s%-3d  %-6s -> %s"
            % (key[0], key[1], have or "(none)", want))
        if have:
            changed += 1
        else:
            added += 1
        if apply_it:
            # Identity first, patch second -- and the identity carries the
            # keys the store matches on, not just the gid, so this updates
            # the record that exists rather than minting a second one.
            store.upsert_session(
                {"gid": rec.get("gid"), "key": rec.get("key"),
                 "loose_key": rec.get("loose_key"),
                 "mouse": key[0], "session": key[1],
                 "label": rec.get("label")},
                {"hemisphere": want,
                 "hemisphere_source":
                     "EEGData_FeederSheet_PTEN_with_DKO.xlsx"})
    say("\n  unchanged %d   newly set %d   replaced %d   skipped %d"
        % (same, added, changed, skipped))
    return {"same": same, "added": added, "changed": changed,
            "skipped": skipped}


def do_layers(store, layers, index, feed, apply_it):
    say("\n" + "=" * 70)
    say("LAYERS  (v0 unlabelled, v1 imported)")
    say("=" * 70)
    made = same = skipped = unknown = 0
    unknown_locs = collections.Counter()
    for key in sorted(feed):
        chans = feed[key]["chans"]
        if not chans:
            continue
        recs = index.get(key) or []
        if len(recs) != 1:
            say("  m%-3d s%-3d  SKIPPED: %d recordings" % (key[0], key[1], len(recs)))
            skipped += 1
            continue
        rec = recs[0]
        gid = rec.get("gid")

        mapping, bad_locs = {}, []
        for ch, loc in sorted(chans.items()):
            region = LOCATION_TO_REGION.get(loc.upper())
            if region is None:
                bad_locs.append(loc)
                unknown_locs[loc] += 1
                continue
            mapping[ch] = region
        if bad_locs:
            unknown += 1

        got = layers.get(gid) or {}
        have = got.get("labels") or {}
        versions = got.get("versions") or []
        if have and len(versions) >= 2:
            same += 1
            continue

        counts = collections.Counter(mapping.values())
        say("  m%-3d s%-3d  %s  %d channels -> %s"
            % (key[0], key[1], gid, len(mapping),
               ", ".join("%s:%d" % (k, v) for k, v in sorted(counts.items()))))
        if bad_locs:
            say("            unmapped locations: %s"
                % sorted(set(bad_locs)))
        made += 1
        if apply_it:
            layers.import_versions(
                gid, mapping,
                session_label=rec.get("label"),
                note="Migrated from EEGData_FeederSheet_PTEN_with_DKO.xlsx",
                by="feeder sheet import")
    say("\n  would create %d   already versioned %d   skipped %d"
        % (made, same, skipped))
    if unknown_locs:
        say("  locations with no region: %s" % dict(unknown_locs))
    return {"made": made, "same": same, "skipped": skipped}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--apply", action="store_true",
                    help="actually write. Without this it only reports.")
    ap.add_argument("--only", default="bad,side,layers",
                    help="comma-separated: bad, side, layers")
    ap.add_argument("--bad-mode", default="report",
                    choices=("report", "union", "replace"),
                    help="report: say what differs and write nothing "
                         "(default). union: add the sheet's without ever "
                         "removing one BARRY already holds. replace: the "
                         "sheet wins, including un-flagging channels.")
    args = ap.parse_args()
    jobs = {j.strip() for j in args.only.split(",") if j.strip()}

    import backend.app as A

    say("Reading the spreadsheets...")
    bad = read_bad()
    feed, conflicts = read_feeder()
    say("  bad-channel rows: %d" % len(bad))
    say("  feeder sessions:  %d  (%d channel rows)"
        % (len(feed), sum(len(v["chans"]) for v in feed.values())))
    if conflicts:
        say("  !! %d cells where the tabs disagree -- these are SKIPPED:"
            % len(conflicts))
        for key, ch, bysheet in conflicts[:10]:
            say("     m%ds%d ch%d: %s" % (key[0], key[1], ch, bysheet))
    else:
        say("  the eight tabs agree everywhere they overlap")

    index = barry_index(A.STORE)
    say("  BARRY (mouse, session) keys: %d" % len(index))
    if not args.apply:
        say("\n*** DRY RUN -- nothing will be written. Add --apply to do it.")

    if "bad" in jobs:
        do_bad(A.STORE, index, bad, args.apply, mode=args.bad_mode)
    if "side" in jobs:
        do_side(A.STORE, index, feed, args.apply)
    if "layers" in jobs:
        do_layers(A.STORE, A.LAYERS, index, feed, args.apply)

    say("\n" + ("Applied." if args.apply
                else "Nothing was written. Re-run with --apply."))


if __name__ == "__main__":
    main()
