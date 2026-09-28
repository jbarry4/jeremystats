# -*- coding: utf-8 -*-
"""Assert `backend/histo.py` still says what the histology workbook says.

`histo.py` holds the scores as literals, because a module the whole app
imports must not read a spreadsheet at import time -- the same rule
`probes.py` follows, and `tools/check_electrode_map.py` is its sibling.
That is only safe if something re-reads the real file, so this does.

Three failures it is built to catch, all of which draw a perfectly normal
matrix:

  1. A cell changed in the workbook and not here. The matrix then labels a
     probe with a region the histology no longer claims.
  2. A cell re-transcribed and quietly tidied -- "L entorhinal (need second
     opinion)" written down as "entorhinal". The comparison is on the RAW
     string for exactly this reason.
  3. A region label that binds to no workbook column. `_norm` exists
     because three files spell dHC three ways; if it ever stops matching,
     every probe in that region scores `unscored` and the whole thing
     passes vacuously with a clean-looking result.

    python tools\\check_histo.py

Run it from PowerShell.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, APP)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:                                        # noqa: BLE001
    pass

from backend import histo, probes                        # noqa: E402

BOOK = os.path.join(APP, histo.SHEET_FILE)

fails = []
notes = []


def bad(msg):
    fails.append(msg)
    print("FAIL  " + msg)


def ok(msg):
    print("ok    " + msg)


def main():
    if not os.path.isfile(BOOK):
        print("The workbook is not here: " + BOOK)
        print("Nothing was checked. That is a failure, not a skip -- a "
              "check that cannot find its subject must not report a pass.")
        return 1

    try:
        import openpyxl
    except ImportError:
        print("openpyxl is not installed, so the workbook cannot be read.")
        return 1

    wb = openpyxl.load_workbook(BOOK, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[("" if c is None else str(c).strip()) for c in r]
            for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(x for x in r)]
    if not rows:
        print("The workbook has no rows.")
        return 1

    head = rows[0]
    cols = [c for c in head[1:] if c]

    # ---- the columns ----
    if tuple(cols) == tuple(histo.SHEET_COLUMNS):
        ok("the twelve columns are the twelve columns, in order")
    else:
        bad("the columns moved.\n      workbook: %s\n      histo.py: %s"
            % (cols, list(histo.SHEET_COLUMNS)))

    # ---- the cells, raw ----
    seen = {}
    for r in rows[1:]:
        name = (r[0] or "").strip()
        if not name:
            continue
        try:
            rat = int(name.lstrip("Jj"))
        except ValueError:
            bad("row label %r is not a rat" % name)
            continue
        seen[rat] = {c: (r[1 + i] if 1 + i < len(r) else "")
                     for i, c in enumerate(cols)}

    if set(seen) == set(histo.HISTO_RAW):
        ok("the same %d rats are scored: %s"
           % (len(seen), ", ".join("J%d" % n for n in sorted(seen))))
    else:
        bad("the scored rats differ. workbook %s, histo.py %s"
            % (sorted(seen), sorted(histo.HISTO_RAW)))

    diffs = 0
    for rat in sorted(set(seen) & set(histo.HISTO_RAW)):
        for col in cols:
            want = seen[rat].get(col, "")
            got = histo.HISTO_RAW[rat].get(col, "")
            if want != got:
                diffs += 1
                bad("J%d %s: workbook %r, histo.py %r"
                    % (rat, col, want, got))
    if not diffs:
        ok("all %d cells match the workbook character for character"
           % sum(len(v) for v in seen.values()))

    # ---- every cell is understood ----
    #
    # A cell nobody anticipated is not a crash; `read_cell` keeps it and
    # flags it. But it has to be SAID, because until somebody adds it to
    # OTHER_REGIONS the matrix draws that probe under a name typed into a
    # spreadsheet rather than a region this app knows.
    unknown = {}
    for rat, row in sorted(seen.items()):
        for col, raw in row.items():
            got = histo.read_cell(raw)
            if not got.get("recognised"):
                unknown.setdefault(raw, []).append("J%d %s" % (rat, col))
    if unknown:
        for raw, where in sorted(unknown.items()):
            notes.append("unrecognised region %r in %s -- add it to "
                         "OTHER_REGIONS in backend/histo.py"
                         % (raw, ", ".join(where)))
        print("note  %d cell value(s) name a region histo.py does not know"
              % len(unknown))
    else:
        ok("every cell value is one histo.py understands")

    # ---- the binding, which is the one that fails silently ----
    #
    # Three spellings of dHC and two of Prh are in play across
    # electrode_map.xlsx, this workbook and probes.py. If the normaliser
    # stops matching, `probe_sanity` scores that region `unscored` for
    # every rat and reports a full, plausible, entirely empty result.
    groups = probes.regions_for("dewey32", list(range(1, 33)))
    if len(groups) != 12:
        bad("probes.regions_for gave %d groups, not 12" % len(groups))
    missed = [g.get("region") for g in groups
              if not histo.sheet_column(g.get("sheet_label")
                                        or g.get("region"))]
    if missed:
        bad("these probe regions bind to no workbook column, so they would "
            "score `unscored` for every rat: %s" % missed)
    else:
        ok("all 12 probe regions bind to a workbook column")

    bound = sorted(histo.sheet_column(g.get("sheet_label") or g.get("region"))
                   for g in groups)
    if len(set(bound)) != len(bound):
        bad("two probe regions bind to the same column: %s" % bound)
    elif not missed:
        ok("and each binds to a different one")

    # ---- the interpretation, end to end ----
    for rat in sorted(histo.HISTO_RAW):
        recs = histo.probe_sanity(rat, groups)
        if len(recs) != 12:
            bad("J%d produced %d records, not 12" % (rat, len(recs)))
            continue
        if any(r["verdict"] == histo.UNSCORED for r in recs):
            bad("J%d has a scored row but %d of its probes came back "
                "`unscored` -- the binding is not doing what it looks "
                "like it is doing"
                % (rat, sum(1 for r in recs
                            if r["verdict"] == histo.UNSCORED)))
    if not fails:
        ok("every scored rat produces twelve records and none of them "
           "falls through to `unscored`")

    # A rat with no row must come back unscored on all twelve rather than
    # raising or quietly reporting `intended`.
    absent = histo.probe_sanity(5, groups)
    if len(absent) == 12 and all(r["verdict"] == histo.UNSCORED
                                 for r in absent):
        ok("J5, which has images but no row, is `unscored` on all twelve")
    else:
        bad("J5 should be `unscored` on all twelve, got %s"
            % sorted({r["verdict"] for r in absent}))

    print("")
    for n in notes:
        print("note  " + n)
    if fails:
        print("\n%d failure(s)." % len(fails))
        return 1
    print("histo.py agrees with the workbook.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
