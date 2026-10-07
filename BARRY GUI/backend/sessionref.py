# -*- coding: utf-8 -*-
"""sessionref.py -- the lab's session workbook, as a reference.

`PTEN_KCNT1 Dentate Spike Data .xlsx`, in the repo beside this app, is
where the lab keeps what each recording session IS: for PTEN, its condition
(baseline or CNO), its mouse's group (CTL / PTEN / DKO) and subgroup
(CTL / IED+ / IED- / DKO); for KCNT1, the mouse's sex and genotypes. Asked
for 2026-10-06 ("integrate this for our reference"), for the pool's
category chips -- "all CTL", "PTEN IED+", "CNO" -- and its hovers.

READ, NEVER WRITTEN. The workbook is edited by hand while Jarvis runs (it
changed the morning this was written), so it is read live and read again
whenever its modification time or size moves, and nothing here ever writes
to it.

KEYED ON PROJECT, MOUSE AND SESSION, never mouse and session alone: the
numbering restarts per project (memory: "mouse+session is not an identity"),
so PTEN m1 s2 and KCNT1 m1 s2 are two different sessions.

WHERE IT DISAGREES it is said, not settled: `disagreements()` compares each
session with the mouse book's group / subgroup and with the session's own
folder name ("CNO" in it), and returns every mismatch for a person to look
at. The workbook is the reference; the others are how a mistake in it, or
in them, gets noticed.
"""
from __future__ import annotations

import os
import re
import threading

SHEET_FILE = "PTEN_KCNT1 Dentate Spike Data .xlsx"
PTEN_SHEET = "PTEN Recording Sessions"
KCNT1_SHEET = "KCNT1 Recording Sessions"

CONDITION = {"base": "Baseline", "baseline": "Baseline", "cno": "CNO"}

_LOCK = threading.Lock()
_CACHE = {"key": None, "data": None}


def path_of(app_dir):
    return os.path.join(app_dir, SHEET_FILE)


def _int(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else None


def _mouse(v):
    """'m13', 'M13', 13, 13.0 -> 13."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return _int(v)
    m = re.match(r"^\s*[mM]?\s*(\d+)\s*$", str(v))
    return int(m.group(1)) if m else None


def _txt(v):
    if v is None:
        return None
    # A whole number typed into a cell comes back a float: 59, not "59.0".
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip()
    return s or None


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    try:
        head = next(it)
    except StopIteration:
        return [], {}
    idx = {str(h).strip(): j for j, h in enumerate(head) if h is not None}
    return list(it), idx


def _read(path):
    import openpyxl                      # first use only; see lazyimp.py
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        if PTEN_SHEET in wb.sheetnames:
            rows, ix = _rows(wb[PTEN_SHEET])
            g = lambda r, k: (r[ix[k]] if k in ix and ix[k] < len(r)
                              else None)
            for r in rows:
                mouse, sess = _mouse(g(r, "mouse_id")), _int(g(r, "session"))
                if mouse is None or sess is None:
                    continue
                cond = _txt(g(r, "condition"))
                out[("PTEN", mouse, sess)] = {
                    "project": "PTEN", "mouse": mouse, "session": sess,
                    "condition": CONDITION.get((cond or "").lower(), cond),
                    "group": _txt(g(r, "group")),
                    "subgroup": _txt(g(r, "subgroup")),
                    "bad_channels": _txt(g(r, "bad channel")),
                    "curated_by": _txt(g(r, "Cur_Initials")),
                }
        if KCNT1_SHEET in wb.sheetnames:
            rows, ix = _rows(wb[KCNT1_SHEET])
            g = lambda r, k: (r[ix[k]] if k in ix and ix[k] < len(r)
                              else None)
            for r in rows:
                mouse, sess = _mouse(g(r, "Mouse ID")), _int(g(r, "Session"))
                if mouse is None or sess is None:
                    continue
                geno = {k: _txt(g(r, k)) for k in ("YH", "SST", "PV")
                        if _txt(g(r, k)) not in (None, "N/A")}
                out[("KCNT1", mouse, sess)] = {
                    "project": "KCNT1", "mouse": mouse, "session": sess,
                    "condition": None,
                    # The genotype that names the KCNT1 cohort.
                    "group": ("YH " + geno["YH"]) if geno.get("YH") else None,
                    "subgroup": None, "sex": _txt(g(r, "Sex")),
                    "genotype": geno,
                    "bad_channels": _txt(g(r, "Bad Channel")),
                    "curated_by": _txt(g(r, "Cur_Initials")),
                }
    finally:
        try:
            wb.close()
        except Exception:                                # noqa: BLE001
            pass
    return out


def load(app_dir):
    """{(project, mouse, session): facts}, read again when the file moves.

    {} when the workbook is not there: a clone without it pools with the
    registry's own filing, as before.
    """
    p = path_of(app_dir)
    try:
        st = os.stat(p)
    except OSError:
        return {}
    key = (st.st_mtime_ns, st.st_size)
    with _LOCK:
        if _CACHE["key"] == key and _CACHE["data"] is not None:
            return _CACHE["data"]
    try:
        data = _read(p)
    except Exception:                                    # noqa: BLE001
        # Open in Excel and locked, or half-saved: the last good reading
        # stands rather than an empty one.
        with _LOCK:
            return _CACHE["data"] or {}
    with _LOCK:
        _CACHE.update(key=key, data=data)
    return data


def facts_for(app_dir, project, mouse, session):
    """The workbook's word on one session, or None."""
    m, s = _mouse(mouse), _int(session)
    if not project or m is None or s is None:
        return None
    return load(app_dir).get((str(project), m, s))


_CNO = re.compile(r"cno", re.I)


def folder_condition(paths, label=None):
    """What the session's folder (or label) says: 'CNO' when any of it
    names CNO, else None -- a folder without the word says nothing either
    way, since baseline sessions are not always marked."""
    for p in list(paths or []) + ([label] if label else []):
        tail = "\\\\".join(str(p or "").replace("/", "\\\\").split("\\\\")[-3:])
        if _CNO.search(tail):
            return "CNO"
    return None


def disagreements(app_dir, sessions, mice):
    """Every place the workbook and another record of the same thing differ.

    `sessions`: registry rows (project, mouse, session, paths, label, gid).
    `mice`: {(project, mouse): {"group", "subgroup"}} from the mouse book.
    """
    ref = load(app_dir)
    out = []
    seen = set()
    for r in sessions:
        key = (r.get("project"), _mouse(r.get("mouse")), _int(r.get("session")))
        if None in key[1:] or key in seen or key not in ref:
            continue
        seen.add(key)
        f = ref[key]
        mb = mice.get((key[0], key[1])) or {}
        for fld in ("group", "subgroup"):
            if f.get(fld) and mb.get(fld) and str(f[fld]) != str(mb[fld]):
                out.append({"project": key[0], "mouse": key[1],
                            "session": key[2], "field": fld,
                            "workbook": f[fld], "other": mb[fld],
                            "other_source": "mouse book"})
        fc = folder_condition(r.get("paths"), r.get("label"))
        if fc and f.get("condition") and fc != f["condition"]:
            out.append({"project": key[0], "mouse": key[1],
                        "session": key[2], "field": "condition",
                        "workbook": f["condition"], "other": fc,
                        "other_source": "folder name"})
    return out
