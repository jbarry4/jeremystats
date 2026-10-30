"""
finalstage.py -- the Monolith page's twelfth tab: the final stage.

What the lab's two last documents ask for (Shahriar, 2026-10-30), on the
same rats, wires and windows as tab 11 (backend/evenmore.py), from the
same pooled rat-days:

  * Within-region power (Within Region Power Analysis Plan, 2026-10-29;
    Supplemental Notes, 2026-10-30). Each region on its own (RSC, dHC,
    ACC, OFC, PER; the notes: Region is never a factor), each band on its
    own: 2 (Session) x 3 (Period: baseline, cue 1, cue 2), and the same
    across frequency bins. Welch power with Dickson's settings, pooled
    over a session's trials (for power, pooling is the mean: nothing to
    flag), in dB re 1 uV^2/Hz; a region is the mean of its valid sides.
    No per-rat normalisation in the model: the rat is the cluster, and
    Session x Period cancels day-to-day drift. Full 10 s windows: the
    cue-onset response is part of what is being measured (the user).
  * Oscillation against background (specparam / FOOOF): the aperiodic
    exponent and offset, and the theta peak's frequency and power, 2 x 3
    in each region.
  * Which region leads (Granger), on the pairs tab 11 found: spectral
    Granger from each rat-day's pooled cross-spectral matrix (every 1 s
    piece of every trial at 250 Hz, Wilson's factorisation), A->B and
    B->A, Early against Late.
  * Theta-gamma coupling within each region of those pairs: Tort's MI of
    the pooled phase-amplitude histogram, Early against Late.
  * The side question (not for Joe): gamma coherence with every trial its
    own row, nothing pooled or averaged.
  * Leave-one-rat-out and a linear mixed model beside every hit.
  * The methods, written from the run itself: a card, a .txt and a .docx.

THE SWITCH: every test is reported with its raw p and its Benjamini-
Hochberg q. With the correction on (the default), a hit is q < 0.05; off,
p < 0.05. Which pairs and regions get Granger and PAC follows the switch,
so both sides' follow-ups are computed, and every figure is drawn both
ways (fig_X and fig_X_raw).

The GEE, as tab 11's: rat as the cluster, exchangeable; Gamma with a log
link for coherence, Granger and MI, normal with the identity for dB.
Tests of up to 2 df on robust (Mancl-DeRouen) errors, more on the
model's own. Coefficients are reported the way SPSS reports them, the
last level of each factor the reference (the user, 2026-10-30).

    finalstage.work(worker)   -> final.json, the figures and the exports
"""
from __future__ import annotations

import csv
import io
import json
import math
import os
import time
import zipfile

import numpy as np

from . import evenmore as E
from . import joe

SCHEMA = 1
REGIONS = tuple(s for s, _say in E.STRUCTS)
REGION_SAY = dict(E.STRUCTS)
#: one hue a region (Okabe and Ito), apart from tab 11's pair colours' order
REGION_COLORS = {"RSC": "#E69F00", "DHC": "#56B4E9", "ACC": "#009E73", "OFC": "#0072B2", "PER": "#D55E00"}
DB_SAY = "dB re 1 µV²/Hz"
FOOOF_RANGE = (2.0, 55.0)              # below the notch: a peak-free fit there needs no gap
FOOOF_PARAMS = (("exponent", "Aperiodic exponent", ""), ("offset", "Aperiodic offset", "log10 µV²/Hz"),
                ("theta_cf", "Theta peak frequency", "Hz"), ("theta_pw", "Theta peak power", "log10 above background"))
ALPHA = E.ALPHA
MODES = E.MODES
TWENTY = ("“20 × 16 × 8 × 2”, in the meeting notes, is waiting for confirmation from Joe. Read as: a 20 s cue pair × 16 trials a "
          "session × 8 rats × 2 sessions.")


class FinalError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def _root(*parts):
    from . import monolith as MO
    return MO._path("final", *parts)


def figs_dir(*parts):
    return _root("figs", *parts)


def export_dir(*parts):
    return _root("export", *parts)


# --------------------------------------------------------------------------
# Power, from tab 11's pooled rat-days
# --------------------------------------------------------------------------
def _scale():
    win = __import__("backend.coupling", fromlist=["matlab_hanning"]).matlab_hanning(E.SEG)
    return 2.0 / (E.FS * float(np.sum(win * win)))


def side_psd(arrs, w):
    """(sides, NF) one-sided power spectral density in µV²/Hz, each side's
    pooled over every piece of every trial in window `w`; NaN where the
    side was never valid."""
    wi = E.WINS.index(w)
    n = arrs["pseg"][wi].astype(float)
    with np.errstate(invalid="ignore", divide="ignore"):
        out = arrs["psd"][wi] / n[:, None] * _scale()
    out[n == 0] = np.nan
    return out


def region_sides(r):
    """The indices of a region's two sides in E.SIDE_NAMES."""
    return [E.SIDE_NAMES.index("%s %s" % (s, r)) for s in E.SIDES]


def _db(x):
    with np.errstate(invalid="ignore", divide="ignore"):
        return 10.0 * np.log10(x)


def power_tables(days):
    """band[w][band][(rat, day)] -> (5,) region dB; bins[w][band][(rat,
    day)] -> (5, nbins) dB; psd[w][(rat, day)] -> (sides, NF) linear. A
    region is the mean of its valid sides, in dB (each side counting the
    same, as tab 11's wire pairs)."""
    out = {"band": {w: {b[0]: {} for b in E.BANDS} for w in E.PERIODS},
           "bins": {w: {b[0]: {} for b in E.BANDS} for w in E.PERIODS}, "psd": {w: {} for w in E.PERIODS}}
    for k, (arrs, _m) in sorted(days.items()):
        for w in E.PERIODS:
            P = side_psd(arrs, w)
            out["psd"][w][k] = P
            for bid, lo, hi in E.BANDS:
                m = joe.band_mask(E.FREQS, lo, hi)
                side_band = _db(E._nanmean(P[:, m], axis=1))                          # (sides,)
                bins = np.stack([_db(E._nanmean(P[:, mm], axis=1)) for _f, mm in E.bins_of(bid)], axis=1)   # (sides, nbins)
                out["band"][w][bid][k] = np.array([E._nanmean(side_band[region_sides(r)], axis=0) for r in REGIONS])
                out["bins"][w][bid][k] = np.stack([E._nanmean(bins[region_sides(r)], axis=0) for r in REGIONS])
    return out


def region_rats(PT, ri):
    """The rats with region `ri` in every period on both days."""
    rats = sorted({r for r, _d in PT["band"]["base10"]["theta"]})
    out = []
    for r in rats:
        vals = [PT["band"][w]["theta"].get((r, d), np.full(len(REGIONS), np.nan))[ri] for w in E.PERIODS for d in E.DAYS]
        if all(np.isfinite(v) for v in vals):
            out.append(r)
    return out


def power_models(PT, band):
    """For each region: E1 (2 x 3 on band dB; Session x Period and its three
    planned contrasts) and E2 (2 x 3 x F); the tests corrected across the
    five regions."""
    per, rows_all, rows_bins = [], [], []
    for ri, r in enumerate(REGIONS):
        rats = region_rats(PT, ri)
        if len(rats) < E.MIN_RATS:
            per.append(None)
            continue
        rows, rb = [], []
        for rat in rats:
            for d in E.DAYS:
                for w in E.PERIODS:
                    rows.append({"rat": rat, "day": d, "period": E.PERIOD_CODE[w], "region": r, "y": float(PT["band"][w][band][(rat, d)][ri])})
                    for fi, (f, _m) in enumerate(E.bins_of(band)):
                        v = PT["bins"][w][band][(rat, d)][ri, fi]
                        if np.isfinite(v):
                            rb.append({"rat": rat, "day": d, "period": E.PERIOD_CODE[w], "region": r, "f": "f%02d" % fi, "hz": f, "y": float(v)})
        df, dfb = E._frame(rows), E._frame(rb)
        ts, err = E._safe(lambda: E.terms(E.gee(df, "y ~ C(day, Sum) * C(period, Sum)", "normal")))
        ts2, err2 = E._safe(lambda: E.terms(E.gee(dfb, "y ~ C(day, Sum) * C(period, Sum) * C(f, Sum)", "normal")))
        sp = next((t for t in ts or [] if t["term"] == "Session × Period"), None)
        spf = next((t for t in ts2 or [] if t["term"] == "Session × Period × Frequency"), None)
        x = {"region": r, "rats": rats, "e1": ts, "e1_error": err, "e2": ts2, "e2_error": err2,
             "sp": dict(sp) if sp else None, "spf": dict(spf) if spf else None,
             "contrasts": E.period_contrasts(df, "normal"), "rows": rows, "rows_bins": rb}
        per.append(x)
        rows_all += rows
        rows_bins += rb
    E.fdr_across(per, lambda x: (x["sp"] or {}).get("p"), lambda x, q: x["sp"].__setitem__("q", q))
    E.fdr_across(per, lambda x: (x["spf"] or {}).get("p"), lambda x, q: x["spf"].__setitem__("q", q))
    for a, b, _s in E.CONTRASTS:
        k = E.contrast_key(a, b)
        E.fdr_across(per, lambda x: x["contrasts"][k].get("p"), lambda x, q: x["contrasts"][k].__setitem__("q", q))
    return per


# --------------------------------------------------------------------------
# Oscillation against background (FOOOF)
# --------------------------------------------------------------------------
def fooof_fit(freqs, psd):
    """FOOOF on one spectrum (linear power), 2-55 Hz, fixed aperiodic:
    {exponent, offset, theta_cf, theta_pw} (the theta peak the strongest
    one centred in 5-12 Hz; NaN when there is none)."""
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from fooof import FOOOF
        fm = FOOOF(peak_width_limits=(1.0, 8.0), max_n_peaks=6, min_peak_height=0.05, aperiodic_mode="fixed", verbose=False)
        ok = np.isfinite(psd) & (psd > 0)
        if ok.sum() < 20:
            return None
        try:
            fm.fit(freqs[ok], psd[ok], list(FOOOF_RANGE))
        except Exception:                                    # noqa: BLE001
            return None
    if not fm.has_model:
        return None
    off, exp = fm.aperiodic_params_
    th = [pk for pk in fm.peak_params_ if 5.0 <= pk[0] <= 12.0]
    cf, pw = (max(th, key=lambda pk: pk[1])[:2]) if th else (np.nan, np.nan)
    return {"exponent": float(exp), "offset": float(off), "theta_cf": float(cf), "theta_pw": float(pw), "r2": float(fm.r_squared_)}


def fooof_tables(PT):
    """[(rat, day, period, region)] -> parameters, each region the mean of
    its sides' fits."""
    out = {}
    for w in E.PERIODS:
        for (rat, d), P in PT["psd"][w].items():
            fits = [fooof_fit(E.FREQS, P[si]) for si in range(P.shape[0])]
            for r in REGIONS:
                got = [fits[si] for si in region_sides(r) if fits[si]]
                if got:
                    out[(rat, d, w, r)] = {k: float(np.nanmean([g[k] for g in got])) if any(np.isfinite(g[k]) for g in got) else float("nan")
                                           for k in ("exponent", "offset", "theta_cf", "theta_pw", "r2")}
    return out


def fooof_models(FT):
    """Session x Period (and its contrasts) on each parameter in each region,
    corrected across the five regions for each parameter."""
    out = {}
    for key, _say, _u in FOOOF_PARAMS:
        per = []
        for r in REGIONS:
            rows = []
            rats = sorted({k[0] for k in FT if k[3] == r})
            for rat in rats:
                cells = [FT.get((rat, d, w, r), {}).get(key) for d in E.DAYS for w in E.PERIODS]
                if not all(c is not None and np.isfinite(c) for c in cells):
                    continue
                for d in E.DAYS:
                    for w in E.PERIODS:
                        rows.append({"rat": rat, "day": d, "period": E.PERIOD_CODE[w], "region": r, "y": float(FT[(rat, d, w, r)][key])})
            n = len({x["rat"] for x in rows})
            if n < E.MIN_RATS:
                per.append({"region": r, "rats": sorted({x["rat"] for x in rows}), "error": "fewer than %d rats with a fit in all six cells" % E.MIN_RATS,
                            "sp": None, "contrasts": {}, "rows": rows})
                continue
            df = E._frame(rows)
            ts, err = E._safe(lambda: E.terms(E.gee(df, "y ~ C(day, Sum) * C(period, Sum)", "normal")))
            sp = next((t for t in ts or [] if t["term"] == "Session × Period"), None)
            per.append({"region": r, "rats": sorted({x["rat"] for x in rows}), "terms": ts, "error": err, "sp": dict(sp) if sp else None,
                        "contrasts": E.period_contrasts(df, "normal"), "rows": rows})
        E.fdr_across(per, lambda x: (x.get("sp") or {}).get("p"), lambda x, q: x["sp"].__setitem__("q", q))
        for a, b, _s in E.CONTRASTS:
            k = E.contrast_key(a, b)
            E.fdr_across(per, lambda x: (x.get("contrasts") or {}).get(k, {}).get("p"), lambda x, q: x["contrasts"][k].__setitem__("q", q))
        out[key] = per
    return out


# --------------------------------------------------------------------------
# The hits, under each side of the switch
# --------------------------------------------------------------------------
def hits(EM, mode):
    """[(pair, band)] tab 11 found: each pair's Late vs Early (phase 1)
    passing under the switch."""
    out = []
    for B in EM["bands"]:
        for x in B["phase1"]["pairs"]:
            if x and E.sig(x.get("a2"), mode):
                out.append((x["pair"], B["band"]))
    return out


def hit_regions(EM, mode):
    out = []
    for p, _b in hits(EM, mode):
        for i in E.PAIRS[p]:
            if REGIONS[i] not in out:
                out.append(REGIONS[i])
    return [r for r in REGIONS if r in out]


# --------------------------------------------------------------------------
# Granger, from the pooled cross-spectral matrices
# --------------------------------------------------------------------------
GC_FREQS = np.fft.fftfreq(E.GC_NFFT, 1.0 / joe.GC_FS)


def granger_of(days, p, band):
    """{(rat, day): (ab, ba, spec_ab, spec_ba)}: each valid wire pair's
    pooled matrix through Wilson's factorisation, band means and spectra,
    then the mean of the wire pairs."""
    from . import sweep
    _b, lo, hi = next(b for b in E.BANDS if b[0] == band)
    m = joe.band_mask(GC_FREQS, lo, hi) & (GC_FREQS > 0)
    pos = (GC_FREQS > 0) & (GC_FREQS <= 100)
    out = {}
    for (rat, d), (arrs, _meta) in sorted(days.items()):
        n = arrs["gseg"][p]
        valid = [c for c in range(len(E.COMBOS)) if n[c] > 0]
        if not valid:
            continue
        S = np.stack([arrs["gS"][p, c] / n[c] for c in valid])
        try:
            H, sig, _it = sweep.wilson(S)
            g12, g21 = sweep.granger(H, sig)
        except np.linalg.LinAlgError:
            continue
        out[(rat, d)] = (float(np.nanmean(np.nanmean(g12[:, m], axis=1))), float(np.nanmean(np.nanmean(g21[:, m], axis=1))),
                         np.nanmean(g12[:, pos], axis=0), np.nanmean(g21[:, pos], axis=0))
    return out


def granger_models(days, EM):
    """Every pair and band that is a hit on either side of the switch:
    Late vs Early in each direction (Gamma, log link, robust), Session x
    Direction; q across the items that are hits with the correction on."""
    items = []
    for p, band in sorted(set(hits(EM, "raw")) | set(hits(EM, "fdr"))):
        G = granger_of(days, p, band)
        rats = sorted({r for r, d in G if (r, "Precon1") in G and (r, "Precon4") in G})
        rows = []
        for r in rats:
            for d in E.DAYS:
                ab, ba, _s1, _s2 = G[(r, d)]
                if ab > 0:
                    rows.append({"rat": r, "day": d, "dir": "a to b", "y": ab})
                if ba > 0:
                    rows.append({"rat": r, "day": d, "dir": "b to a", "y": ba})
        df = E._frame(rows)
        x = {"pair": p, "band": band, "rats": rats, "a": E.STRUCTS[E.PAIRS[p][0]][1], "b": E.STRUCTS[E.PAIRS[p][1]][1],
             "in": {m: (p, band) in hits(EM, m) for m in MODES}, "rows": rows, "G": G}
        if len(rats) < E.MIN_RATS:
            x["error"] = "fewer than %d rats" % E.MIN_RATS
            items.append(x)
            continue
        for k, dname in (("ab", "a to b"), ("ba", "b to a")):
            sub = df[df["dir"] == dname]
            got, err = E._safe(lambda: E.ratio(E.gee(sub, "y ~ C(day, Treatment('Precon1'))"), "Precon4"))
            x[k] = got or {"error": err}
        sd, err = E._safe(lambda: E.interaction(df, "y ~ C(day, Treatment('Precon1')) * C(dir, Treatment('a to b'))", ["Precon4", "b to a"]))
        x["sd"] = sd or {"error": err}
        items.append(x)
    fam = [x for x in items if x["in"]["fdr"] and "error" not in x]
    for k in ("ab", "ba", "sd"):
        qs = E.bh([x[k].get("p") for x in fam])
        for x, q in zip(fam, qs):
            if q is not None:
                x[k]["q"] = q
    return items


# --------------------------------------------------------------------------
# PAC within each region of the hit pairs
# --------------------------------------------------------------------------
def pac_of(days, r):
    """{(rat, day): (MI, profile)}: each valid side's pooled histogram's
    Tort MI and its normalised amplitude-by-phase profile, the mean of the
    region's sides."""
    out = {}
    N = E.PAC_BINS
    for (rat, d), (arrs, _m) in sorted(days.items()):
        mis, profs = [], []
        for si in region_sides(r):
            n = arrs["pac_n"][si]
            if n.sum() <= 0:
                continue
            with np.errstate(invalid="ignore", divide="ignore"):
                mean = arrs["pac_sum"][si] / n
                pr = mean / np.nansum(mean)
                h = -np.nansum(np.where(pr > 0, pr * np.log(pr), 0.0))
            mis.append((math.log(N) - h) / math.log(N))
            profs.append(pr)
        if mis:
            out[(rat, d)] = (float(np.mean(mis)), np.mean(profs, axis=0))
    return out


def pac_models(days, EM):
    items = []
    for r in [x for x in REGIONS if x in set(hit_regions(EM, "raw")) | set(hit_regions(EM, "fdr"))]:
        Pm = pac_of(days, r)
        rats = sorted({k[0] for k in Pm if (k[0], "Precon1") in Pm and (k[0], "Precon4") in Pm})
        rows = [{"rat": rat, "day": d, "y": Pm[(rat, d)][0]} for rat in rats for d in E.DAYS]
        x = {"region": r, "rats": rats, "in": {m: r in hit_regions(EM, m) for m in MODES}, "rows": rows, "P": Pm}
        if len(rats) < E.MIN_RATS:
            x["error"] = "fewer than %d rats" % E.MIN_RATS
        else:
            got, err = E._safe(lambda: E.ratio(E.gee(E._frame(rows), "y ~ C(day, Treatment('Precon1'))"), "Precon4"))
            x["le"] = got or {"error": err}
        items.append(x)
    fam = [x for x in items if x["in"]["fdr"] and "le" in x]
    qs = E.bh([x["le"].get("p") for x in fam])
    for x, q in zip(fam, qs):
        if q is not None:
            x["le"]["q"] = q
    return items


# --------------------------------------------------------------------------
# The side question: every trial its own row (not for Joe)
# --------------------------------------------------------------------------
def side_question(days, T):
    """Gamma coherence, each trial's own (its own spectra, the mean of its
    valid wire pairs), every trial a GEE row clustered by rat: Late vs
    Early in each pair, beside tab 11's pooled answer."""
    out = []
    gi = [b[0] for b in E.BANDS].index("gamma")
    for p in range(len(E.PAIRS)):
        rows = []
        for (rat, d), (arrs, meta) in sorted(days.items()):
            tb = arrs["tband"][0][:, p, :, gi]                     # (T, C): the cue pair
            for t in range(tb.shape[0]):
                v = E._nanmean(tb[t], axis=0)
                if np.isfinite(v) and v > 0:
                    rows.append({"rat": rat, "day": d, "trial": t, "y": float(v)})
        df = E._frame(rows)
        rats = sorted({x["rat"] for x in rows if any(y["rat"] == x["rat"] and y["day"] != x["day"] for y in rows)})
        df = df[df["rat"].isin(rats)] if len(df) else df
        if len(rats) < E.MIN_RATS:
            out.append({"pair": p, "rats": rats, "n_rows": len(df), "error": "fewer than %d rats" % E.MIN_RATS})
            continue
        got, err = E._safe(lambda: E.ratio(E.gee(df, "y ~ C(day, Treatment('Precon1'))"), "Precon4"))
        res = E.gee(df, "y ~ C(day, Treatment('Precon1'))") if got else None
        if got and res is not None:
            i = list(res.params.index).index(next(n for n in res.params.index if "Precon4" in n))
            se = float(np.sqrt(np.asarray(res.cov_robust_bc)[i, i]))
            got["lo"], got["hi"] = math.exp(float(res.params.iloc[i]) - 1.959964 * se), math.exp(float(res.params.iloc[i]) + 1.959964 * se)
        out.append({"pair": p, "rats": rats, "n_rows": int(len(df)), "le": got or {"error": err}, "rows": rows})
    E.fdr_across(out, lambda x: (x.get("le") or {}).get("p"), lambda x, q: x["le"].__setitem__("q", q))
    return out


# --------------------------------------------------------------------------
# SPSS's coefficients (the last level the reference)
# --------------------------------------------------------------------------
def spss_coefficients(df, factors, fam):
    """Each coefficient of the full factorial with SPSS's coding: every
    factor's last level (in ascending order) the reference. Robust
    (Mancl-DeRouen) errors and a Wald chi-square on 1 df each."""
    refs = {f: sorted(df[f].unique())[-1] for f in factors}
    rhs = " * ".join("C(%s, Treatment(%r))" % (f, refs[f]) for f in factors)
    res = E.gee(df, "y ~ " + rhs, fam)
    se = np.sqrt(np.diag(np.asarray(res.cov_robust_bc)))
    out = []
    for i, n in enumerate(res.params.index):
        b = float(res.params.iloc[i])
        x = (b / se[i]) ** 2 if se[i] > 0 else float("nan")
        out.append({"term": n, "b": b, "se": float(se[i]), "chi2": x, "p": E._chi2_sf(x, 1) if np.isfinite(x) else None})
    return refs, out


# --------------------------------------------------------------------------
# Drawing
# --------------------------------------------------------------------------
def _save(fig, name):
    os.makedirs(figs_dir(), exist_ok=True)
    out = {}
    for ext in ("svg", "pdf", "png"):
        p = figs_dir("%s.%s" % (name, ext))
        fig.savefig(p, dpi=600 if ext == "png" else None)
        out[ext] = os.path.basename(p)
    return out


def _prism(stem, tables, what):
    from . import joefigs
    return joefigs.prism_files(stem, tables, what, into=figs_dir())


def _sigs(x, mode):
    c = x.get("contrasts") or {}
    return (E.sig(x.get("sp"), mode), E.sig(c.get("cue1-base10"), mode), E.sig(c.get("cue2-base10"), mode))


def _power_lines(x, mode):
    out = E._contrast_line(x, mode)
    spf = x.get("spf")
    if spf:
        out.append("Session × Period × Frequency: χ²(%d) = %s, %s, q = %s (model-based)" % (spf["df"], E.num(spf["chi2"], "%.1f"),
                                                                                         E.p_say(spf["p"]), E.num(spf.get("q"), "%.3f")))
    return [s.replace("×", "×").replace(" dB", " dB") for s in out]


def fig_power_region(PT, band, x, fid, mode):
    """One region, one band: the plan's three panels."""
    plt = E._plt()
    W, H = 183.0, 76.0
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    ri = REGIONS.index(x["region"])
    sym = "θ" if band == "theta" else "γ"
    fig.text(4 / W, 1 - 3 / H, "%s · %s power · baseline, cue 1, cue 2 · %d rats" % (REGION_SAY[x["region"]], E._band_say(band), len(x["rats"])),
             fontsize=7, fontweight="bold", va="top")
    fig.text(4 / W, 1 - 8 / H, "Welch power (Dickson's settings) pooled over each session's trials, both sides averaged, full 10 s windows; "
             "Precon1 solid, Precon4 dashed, mean ± s.e.m. * passes: %s across the five regions." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05, uncorrected"),
             fontsize=5, va="top", color="#4D4D4D")
    E.period_panels(fig, None, band, "power", x["rats"], DB_SAY, REGION_COLORS[x["region"]], "-", 18, 16, H, W,
                    "%s %s power (dB)" % (REGION_SAY[x["region"]], sym), "Change from baseline (dB)", _sigs(x, mode),
                    get=lambda r, d, w: PT["band"][w][band][(r, d)][ri], bins_get=lambda r, d, w: PT["bins"][w][band][(r, d)][ri])
    for k, s in enumerate(_power_lines(x, mode)):
        fig.text(4 / W, 1 - (H - 19 + 3.5 * k) / H, s, fontsize=5.0, va="top", fontweight="bold" if s.endswith("*") else "normal")
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


def fig_power_all(PT, band, per, fid, mode):
    """Every region, a row each, the plan's example laid out in full."""
    plt = E._plt()
    keep = [x for x in per if x]
    W, H = 183.0, 14.0 + 50.0 * len(keep)
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    sym = "θ" if band == "theta" else "γ"
    fig.text(4 / W, 1 - 3 / H, "Within-region %s power · every region · Precon1 solid, Precon4 dashed" % E._band_say(band), fontsize=7.5,
             fontweight="bold", va="top")
    fig.text(4 / W, 1 - 7.5 / H, "dB re 1 µV²/Hz, mean ± s.e.m. of the rats. * Session × Period (a), or Session × (that cue vs baseline) (b, c), "
             "passes: %s across the five regions." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05, uncorrected"), fontsize=5, va="top", color="#4D4D4D")
    for row, x in enumerate(keep):
        ri = REGIONS.index(x["region"])
        y0 = 14.0 + 50.0 * row
        E.period_panels(fig, None, band, "power", x["rats"], DB_SAY, REGION_COLORS[x["region"]], "-", 18, y0, H, W,
                        "%s %s power (dB)" % (REGION_SAY[x["region"]], sym), "Change from baseline (dB)", _sigs(x, mode),
                        letters="abc" if row == 0 else "   ",
                        get=lambda r, d, w, ri=ri: PT["band"][w][band][(r, d)][ri], bins_get=lambda r, d, w, ri=ri: PT["bins"][w][band][(r, d)][ri])
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


def fig_fooof(FM, FT, fid, mode):
    """Four parameters (rows) by five regions (columns), each by period."""
    plt = E._plt()
    W, H = 183.0, 150.0
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    fig.text(4 / W, 1 - 3 / H, "Oscillation against background (FOOOF, 2–55 Hz) · Precon1 solid, Precon4 dashed", fontsize=7.5,
             fontweight="bold", va="top")
    fig.text(4 / W, 1 - 7.5 / H, "Each region's sides fitted separately, then averaged; mean ± s.e.m. of the rats with a fit in all six cells. "
             "* Session × Period passes: %s across the five regions." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05, uncorrected"),
             fontsize=5, va="top", color="#4D4D4D")
    for pi, (key, say, unit) in enumerate(FOOOF_PARAMS):
        for ri, r in enumerate(REGIONS):
            ax = E._ax(fig, 18 + 33 * ri, 16 + 33 * pi, 27, 24, H, W)
            x = next(z for z in FM[key] if z["region"] == r)
            rats = x["rats"]
            if not rats or x.get("error"):
                ax.text(0.5, 0.5, "too few fits", transform=ax.transAxes, ha="center", va="center", fontsize=5, color="#8C8C8C")
                ax.set_xticks([])
                ax.set_yticks([])
            else:
                for di, d in enumerate(E.DAYS):
                    per = np.array([[FT[(rat, d, w, r)][key] for w in E.PERIODS] for rat in rats], float)
                    ax.errorbar([0, 1, 2], [E._mean(per[:, k]) for k in range(3)], yerr=[E._se(per[:, k]) for k in range(3)],
                                color=REGION_COLORS[r], ls="-" if di == 0 else "--", lw=0.9, marker="o", ms=2,
                                mfc=REGION_COLORS[r] if di == 0 else "white", mew=0.5, capsize=1, elinewidth=0.45)
                ax.set_xticks([0, 1, 2])
                ax.set_xticklabels(["B", "C1", "C2"] if pi == 3 else [], fontsize=4.8)
                ax.set_xlim(-0.4, 2.4)
                ax.tick_params(labelsize=4.6)
                if E.sig(x.get("sp"), mode):
                    ax.text(0.97, 0.97, "*", transform=ax.transAxes, ha="right", va="top", fontsize=8, fontweight="bold", color=REGION_COLORS[r])
            if pi == 0:
                ax.set_title(REGION_SAY[r], fontsize=6, pad=2)
            if ri == 0:
                ax.set_ylabel("%s%s" % (say, ("\n(%s)" % unit) if unit else ""), fontsize=5, labelpad=1)
    fig.text(4 / W, 1 - (H - 4) / H, "B baseline, C1 cue 1, C2 cue 2.", fontsize=5, va="bottom", color="#4D4D4D")
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


def _days_panel(ax, vals, col, title, ylab=None, star=False):
    """Each rat Early to Late in grey, the mean ± s.e.m. in colour."""
    for row in vals:
        ax.plot([0, 1], row, color="#BFBFBF", lw=0.4)
    mu = [E._mean(vals[:, k]) for k in range(2)]
    se = [E._se(vals[:, k]) for k in range(2)]
    ax.errorbar([0, 1], mu, yerr=se, color=col, lw=1.1, marker="o", ms=2.6, capsize=1.4, elinewidth=0.55)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(["Early", "Late"])
    ax.set_xlim(-0.35, 1.35)
    ax.set_title(title + ("  *" if star else ""), fontsize=6, pad=2)
    if ylab:
        ax.set_ylabel(ylab, fontsize=5.5, labelpad=1)


def fig_granger(x, fid, mode):
    plt = E._plt()
    W, H = 183.0, 66.0
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    col = E.COLORS[x["pair"]]
    fig.text(4 / W, 1 - 3 / H, "Which leads: %s · %s · %d rats" % (E.PAIR_SAY[x["pair"]], E._band_say(x["band"]), len(x["rats"])),
             fontsize=7, fontweight="bold", va="top")
    fig.text(4 / W, 1 - 8 / H, "Spectral Granger from each rat-day's pooled cross-spectral matrix (cue pair, 250 Hz, 1 s pieces, Wilson's "
             "factorisation), the mean of the valid wire pairs. * Late vs Early passes: %s." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05, uncorrected"),
             fontsize=5, va="top", color="#4D4D4D")
    G = x["G"]
    rats = x["rats"]
    for k, (key, say) in enumerate((("ab", "%s → %s" % (x["a"], x["b"])), ("ba", "%s → %s" % (x["b"], x["a"])))):
        ax = E._ax(fig, 16 + 40 * k, 16, 28, 34, H, W)
        vals = np.array([[G[(r, d)][k] for d in E.DAYS] for r in rats], float)
        _days_panel(ax, vals, col, say, "Granger (%s band)" % x["band"] if k == 0 else None, E.sig(x.get(key), mode))
        E._letter(fig, 7 + 40 * k, 15, "ab"[k], H, W)
    ax = E._ax(fig, 104, 16, 74, 34, H, W)
    f = GC_FREQS[(GC_FREQS > 0) & (GC_FREQS <= 100)]
    for k, ls_dir in ((2, "-"), (3, ":")):
        for di, d in enumerate(E.DAYS):
            S = np.array([G[(r, d)][k] for r in rats], float)
            ax.plot(f, E._nanmean(S, axis=0), color=col if di == 1 else "#8C8C8C", ls=ls_dir, lw=0.9,
                    label="%s, %s" % ("%s → %s" % ((x["a"], x["b"]) if k == 2 else (x["b"], x["a"])), "Late" if di else "Early"))
    _b, lo, hi = next(b for b in E.BANDS if b[0] == x["band"])
    ax.axvspan(lo, hi, color="#EFEFEF", zorder=0)
    ax.set_xlabel("Frequency (Hz)", fontsize=5.5, labelpad=1)
    ax.set_ylabel("Granger, mean of rats", fontsize=5.5, labelpad=1)
    ax.legend(frameon=False, fontsize=4.6, ncol=2, loc="upper right")
    E._letter(fig, 96, 15, "c", H, W)
    lines = []
    for key, say in (("ab", "%s → %s" % (x["a"], x["b"])), ("ba", "%s → %s" % (x["b"], x["a"]))):
        r = x.get(key) or {}
        if r.get("p") is not None:
            lines.append("%s, Late vs Early: ×%s, χ²(1) = %s, %s, q = %s%s" % (say, E.num(r["ratio"]), E.num(r["chi2"], "%.2f"), E.p_say(r["p"]),
                                                                          E.num(r.get("q"), "%.3f"), " *" if E.sig(r, mode) else ""))
    sd = x.get("sd") or {}
    if sd.get("p") is not None:
        lines.append("Session × Direction: ×%s, χ²(1) = %s, %s, q = %s%s" % (E.num(sd["ratio"]), E.num(sd["chi2"], "%.2f"), E.p_say(sd["p"]),
                                                                          E.num(sd.get("q"), "%.3f"), " *" if E.sig(sd, mode) else ""))
    for k, s in enumerate(lines):
        fig.text(4 / W, 1 - (H - 11 + 3.6 * k) / H, s, fontsize=5.1, va="top", fontweight="bold" if s.endswith("*") else "normal")
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


def fig_pac(x, fid, mode):
    plt = E._plt()
    W, H = 120.0, 62.0
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    col = REGION_COLORS[x["region"]]
    fig.text(4 / W, 1 - 3 / H, "Theta–gamma coupling in %s · %d rats" % (REGION_SAY[x["region"]], len(x["rats"])), fontsize=7,
             fontweight="bold", va="top")
    fig.text(4 / W, 1 - 8 / H, "Tort's MI of the pooled histogram: 5–12 Hz phase, 30–90 Hz amplitude, the same wire, the cue pair. "
             "* passes: %s." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05, uncorrected"), fontsize=5, va="top", color="#4D4D4D")
    P = x["P"]
    ax = E._ax(fig, 16, 16, 30, 32, H, W)
    vals = np.array([[P[(r, d)][0] for d in E.DAYS] for r in x["rats"]], float)
    _days_panel(ax, vals, col, "Modulation index", "MI", E.sig(x.get("le"), mode))
    E._letter(fig, 6, 15, "a", H, W)
    ax = E._ax(fig, 62, 16, 54, 32, H, W)
    centres = np.linspace(-180, 180, E.PAC_BINS + 1)[:-1] + 180.0 / E.PAC_BINS
    for di, d in enumerate(E.DAYS):
        prof = E._nanmean(np.array([P[(r, d)][1] for r in x["rats"]]), axis=0)
        ax.plot(centres, prof, color=col, ls="-" if di == 0 else "--", lw=1.0, label="Early" if di == 0 else "Late")
    ax.axhline(1.0 / E.PAC_BINS, color="#8C8C8C", lw=0.4)
    ax.set_xticks([-180, -90, 0, 90, 180])
    ax.set_xlabel("Theta phase (degrees)", fontsize=5.5, labelpad=1)
    ax.set_ylabel("Gamma amplitude (share)", fontsize=5.5, labelpad=1)
    ax.legend(frameon=False, fontsize=5)
    E._letter(fig, 54, 15, "b", H, W)
    le = x.get("le") or {}
    if le.get("p") is not None:
        fig.text(4 / W, 1 - (H - 6) / H, "Late vs Early: ×%s, χ²(1) = %s, %s, q = %s%s" % (E.num(le["ratio"]), E.num(le["chi2"], "%.2f"), E.p_say(le["p"]),
                                                                                       E.num(le.get("q"), "%.3f"), " *" if E.sig(le, mode) else ""),
                 fontsize=5.2, va="top")
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


def fig_side(side, EM, fid, mode):
    """Every trial its own row against tab 11's pooled answer, gamma,
    Late vs Early in each pair."""
    plt = E._plt()
    W, H = 120.0, 80.0
    fig = plt.figure(figsize=(W / 25.4, H / 25.4))
    fig.text(4 / W, 1 - 3 / H, "Side question (not for Joe): gamma, every trial its own row", fontsize=7, fontweight="bold", va="top")
    fig.text(4 / W, 1 - 8 / H, "Late ÷ Early, GEE (Gamma, log link, rat as the cluster) with 95%% interval. Circles: each trial's own "
             "coherence a row. Diamonds: tab 11's pooled answer. * passes: %s." % ("FDR q < 0.05" if mode == "fdr" else "p < 0.05"),
             fontsize=5, va="top", color="#4D4D4D")
    ax = E._ax(fig, 30, 16, 80, 56, H, W)
    gB = next(B for B in EM["bands"] if B["band"] == "gamma")
    for row, x in enumerate(side):
        p = x["pair"]
        st = E._style(p)
        le = x.get("le") or {}
        if le.get("ratio") is not None:
            ax.plot([le["lo"], le["hi"]], [row - 0.12, row - 0.12], color=st["color"], lw=0.9)
            ax.plot([le["ratio"]], [row - 0.12], marker="o", ms=3, color=st["color"], mfc=st["mfc"], mew=0.6)
            if E.sig(le, mode):
                ax.text(le["hi"], row - 0.12, " *", va="center", fontsize=7, fontweight="bold", color=st["color"])
        a2 = ((gB["phase1"]["pairs"][p] or {}).get("a2") or {})
        if a2.get("ratio") is not None:
            ax.plot([a2["ratio"]], [row + 0.18], marker="D", ms=2.6, color="#1A1A1A", mfc="white", mew=0.6)
    ax.axvline(1.0, color="#1A1A1A", lw=0.5)
    ax.set_xscale("log")
    E.ratio_axis(ax)
    ax.set_ylim(len(side) - 0.5, -0.5)
    ax.set_yticks(range(len(side)))
    ax.set_yticklabels([E.PAIR_SAY[x["pair"]] for x in side])
    ax.set_xlabel("Late ÷ Early", fontsize=5.5, labelpad=1)
    files = _save(fig, "fig_%s" % fid)
    plt.close(fig)
    return files


# --------------------------------------------------------------------------
# The methods, from the run itself
# --------------------------------------------------------------------------
def _sampling_rate():
    from . import monolith as MO, nlx
    man = MO.manifest() or {}
    for d in man.get("days") or []:
        for f in d.get("folders") or []:
            if f.get("role") == "SPC" and f.get("local") and os.path.isdir(f["local"]):
                try:
                    hdr = nlx.read_header(os.path.join(f["local"], "CSC1.ncs"))
                    return float(nlx._header_float(hdr, "SamplingFrequency") or 0) or None
                except Exception:                            # noqa: BLE001
                    return None
    return None


def _res(label, x, kind="ratio"):
    if not x or x.get("p") is None:
        return None
    eff = ("×%s" % E.num(x["ratio"])) if x.get("ratio") is not None else ("%s dB" % E.num(x.get("diff"), "%+.2f")) if x.get("diff") is not None else ""
    return {"label": label, "effect": eff, "chi2": x["chi2"], "df": x.get("df", 1), "p": x["p"], "q": x.get("q")}


def significant(EM, F, mode):
    """Every test that passes under the switch, in the methods' words."""
    out = []
    for B in EM["bands"]:
        for x in B["phase1"]["pairs"]:
            if x and E.sig(x.get("a2"), mode):
                out.append(_res("Coherence, %s, %s: Late vs Early" % (E.PAIR_SAY[x["pair"]], B["band"]), x["a2"]))
        for x in B["phase2"]["pairs"]:
            if x and E.sig({"p": E._term_p(x["terms"], "Session × Cue"), "q": x.get("q")}, mode):
                t = next(t for t in x["terms"] if t["term"] == "Session × Cue")
                out.append(_res("Coherence, %s, %s: Session × Cue" % (E.PAIR_SAY[x["pair"]], B["band"]), dict(t, q=x.get("q"))))
            for d, r in (x or {}).get("days", {}).items() if x else []:
                if E.sig(r, mode):
                    out.append(_res("Coherence, %s, %s: cue 2 vs cue 1 on %s" % (E.PAIR_SAY[x["pair"]], B["band"], d), r))
        for x in B["phase3"]["pairs"]:
            if not x:
                continue
            if E.sig(x.get("sp"), mode):
                out.append(_res("Coherence, %s, %s: Session × Period" % (E.PAIR_SAY[x["pair"]], B["band"]), x["sp"]))
            for a, b, say in E.CONTRASTS:
                c = x["contrasts"].get(E.contrast_key(a, b)) or {}
                if E.sig(c, mode):
                    out.append(_res("Coherence, %s, %s: Session × (%s)" % (E.PAIR_SAY[x["pair"]], B["band"], say.lower()), c))
    for band, per in F["power"].items():
        for x in per:
            if not x:
                continue
            if E.sig(x.get("sp"), mode):
                out.append(_res("Power, %s, %s: Session × Period" % (REGION_SAY[x["region"]], band), x["sp"]))
            for a, b, say in E.CONTRASTS:
                c = x["contrasts"].get(E.contrast_key(a, b)) or {}
                if E.sig(c, mode):
                    out.append(_res("Power, %s, %s: Session × (%s)" % (REGION_SAY[x["region"]], band, say.lower()), c))
    for key, say, _u in FOOOF_PARAMS:
        for x in F["fooof"][key]:
            if E.sig(x.get("sp"), mode):
                out.append(_res("%s, %s: Session × Period" % (say, REGION_SAY[x["region"]]), x["sp"]))
    for x in F["granger"]:
        if not x["in"][mode]:
            continue
        for key, say in (("ab", "%s → %s" % (x["a"], x["b"])), ("ba", "%s → %s" % (x["b"], x["a"]))):
            if E.sig(x.get(key), mode):
                out.append(_res("Granger %s, %s: Late vs Early" % (say, x["band"]), x[key]))
        if E.sig(x.get("sd"), mode):
            out.append(_res("Granger, %s, %s: Session × Direction" % (E.PAIR_SAY[x["pair"]], x["band"]), x["sd"]))
    for x in F["pac"]:
        if x["in"][mode] and E.sig(x.get("le"), mode):
            out.append(_res("Theta–gamma MI in %s: Late vs Early" % REGION_SAY[x["region"]], x["le"]))
    return [x for x in out if x]


def methods(EM, F, fs_rec):
    """The methods as sections of paragraphs: the card, the .txt and the
    .docx are all made from this."""
    rats = EM["rats"]
    n_tr = EM.get("trials") or {}
    tr = sorted(n_tr.values())
    S = []
    S.append(("Waiting for confirmation", [TWENTY,
              "Coherence here comes from pooled spectra (every 1 s piece of every trial in a session in one cross-spectrum), not the mean "
              "of per-trial coherence as in Dickson et al. (2022). Mean over rats and pairs, cue pair: " + "; ".join(
                  "%s %s pooled %.3f against %.3f as the mean of trials" % (b, d, EM["flag"][b][d]["pooled"], EM["flag"][b][d]["trialmean"])
                  for b in ("theta", "gamma") for d in E.DAYS) + "."]))
    S.append(("Data included", [
        "%d rats (%s), analysed as one cohort with no between-subject factor. Two sessions: Precon1 (Early) and Precon4 (Late). "
        "Trials per session after the cue pairs were found: %d to %d (a trial is one cue pair, cue 1 then cue 2, each about 10 s)."
        % (len(rats), ", ".join("J%d" % r for r in rats), tr[0] if tr else 0, tr[-1] if tr else 0),
        "Regions: retrosplenial cortex (RSC), dorsal hippocampus (dHC), anterior cingulate (ACC), orbitofrontal cortex (OFC) and "
        "perirhinal cortex (PER), each recorded on both sides; the postrhinal probes were not in postrhinal cortex and are not used. A "
        "side is used only where histology (v2, 2026-10-05) scored its probe “y”. Each side has 2 wires (PER 4); in each "
        "window of each trial the lowest-numbered wire that is not marked bad and not clipped is used (one wire a side).",
        "Windows, aligned to the cue pulses (TTL events; the pairing of two different cue pulses 10 ± 0.25 s apart, TTL 126 the end of "
        "cue 2): the 10 s before cue 1 (baseline), cue 1 (onset to cue 2 onset), cue 2 (onset to its end), and the cue pair (cue 1 + "
        "cue 2, about 20 s). Full windows; the cue-onset response is part of what is measured.",
        "Ten region pairs (RSC–dHC, RSC–ACC, RSC–OFC, RSC–PER, dHC–ACC, dHC–OFC, dHC–PER, ACC–OFC, ACC–PER, OFC–PER). Within a rat, a "
        "region pair's value is the mean of its valid wire pairs, ipsilateral (left–left, right–right) and contralateral (left–right, "
        "right–left) alike; homotopic pairs (one structure, left with right) are never used. A region's power is the mean of its valid "
        "sides, in dB. A pair or region enters a model only with at least %d rats." % E.MIN_RATS]))
    S.append(("Preprocessing", [
        "Neuralynx recordings%s, read in µV (ADBitVolts × 10⁶), decimated to 1 kHz (scipy.signal.decimate, an order-8 Chebyshev "
        "anti-aliasing filter run forwards and backwards). No re-referencing: signals as recorded against the acquisition reference."
        % (" sampled at %g Hz" % fs_rec if fs_rec else ""),
        "Artifact rejection by clipping: a wire is excluded from a window when it sits within 0.5% of the digitiser's rail (ADMaxValue, "
        "±2000 µV input range) for 16 or more consecutive samples, covering 0.5% of the window or one unbroken 50 ms; the next wire of "
        "that side is used, and if none is left the side is missing in that window. Wires marked bad by hand are excluded everywhere.",
        "Line noise: a second-order Butterworth band-stop at 59–61 Hz (MATLAB designfilt, run with filtfilt, as Dickson et al.); the "
        "59–61 Hz bins are left out of every gamma average."]))
    S.append(("Spectral method", [
        "Coherence (magnitude-squared): MATLAB mscohere's settings — a 1024-sample hanning window (MATLAB's, without zero end points), "
        "512 samples of overlap and a 2048-point FFT at 1 kHz (0.49 Hz resolution) — with the cross- and auto-spectra summed over every "
        "piece of every trial of a session before the coherence is taken (pooled). Bands: theta 5–12 Hz; gamma 30–90 Hz without 59–61 Hz. "
        "Frequency bins: theta at the FFT's own 0.49 Hz; gamma in 2 Hz bins.",
        "Power: the same pieces' one-sided power spectral density (µV²/Hz), the mean over every piece of a session's trials, a band's "
        "mean then 10·log10 (dB re 1 µV²/Hz). No normalisation per rat: the rat is the GEE's cluster, and Session × Period cancels "
        "drift between the days.",
        "Oscillation against background: FOOOF 1.1 (specparam), 2–55 Hz, fixed aperiodic mode, peak widths 1–8 Hz, up to 6 peaks; the "
        "aperiodic exponent and offset, and the strongest peak centred in 5–12 Hz (its frequency and power above the background).",
        "Granger causality: each signal at 250 Hz, 1 s pieces half overlapping, mean removed, MATLAB's hanning, a 500-point FFT; the 2 × 2 "
        "cross-spectral matrix pooled over every piece of a session's cue pairs; Wilson's spectral factorisation; spectral Granger "
        "(Geweke) both ways, its band mean; a pair the mean of its valid wire pairs.",
        "Theta–gamma coupling: within one wire, 5–12 Hz phase and 30–90 Hz amplitude (zero-phase band-pass, Hilbert transform, mirrored "
        "padding), 18 phase bins, the amplitude summed by bin over every trial's cue pair, then Tort's modulation index; a region the "
        "mean of its sides."]))
    S.append(("Statistical model", [
        "Every test is a generalised estimating equation (statsmodels 0.15 GEE): the rat as the subject (cluster), an exchangeable "
        "working correlation. Coherence, Granger and the modulation index: Gamma distribution with a log link (effects are ratios). "
        "Power (dB) and the FOOOF parameters: normal distribution with the identity link (effects are differences).",
        "Tests of up to %d degrees of freedom use robust (sandwich) errors with Mancl and DeRouen's small-sample correction; tests of "
        "more (Pair, and any term with Frequency) cannot be estimated robustly from %d clusters and use the model-based errors, which "
        "rest on the working correlation; those are reported but not relied on. Wald chi-square tests." % (E.ROBUST_MAX_DF, len(rats)),
        "Coherence across sessions (each pair): Late vs Early on the band; Session × Frequency across bins. Phase 2 (each pair): "
        "Session × Cue (cue 1, cue 2; every Precon trial is a matched pairing, A→B or C→D, pooled). Coherence and power in the three "
        "periods (each pair; each region on its own, Region is never a factor): Session × Period (2 df), and the planned contrasts "
        "Session × (cue 1 vs baseline), Session × (cue 2 vs baseline), Session × (cue 2 vs cue 1), each on 1 df; power also Session × "
        "Period × Frequency. Granger (each pair and band that passed): Late vs Early in each direction and Session × Direction. "
        "Theta–gamma MI (each region of those pairs): Late vs Early.",
        "Reference category: the tests do not depend on it (sum coding). Coefficients are reported as SPSS reports them, with the last "
        "level of each factor (in ascending order) the reference: Session, Precon4; Period, cue 2; Frequency, the highest bin (11.96 Hz "
        "in theta, 89 Hz in gamma); Direction, the second. The coefficient tables are in the exports.",
        "Correction: Benjamini–Hochberg (FDR) across the ten pairs for each coherence test, across the five regions for each power and "
        "FOOOF test, and within each follow-up family (Granger, MI). The results are reported both corrected (q < 0.05) and uncorrected "
        "(p < 0.05); which pairs are followed up depends on which is used.",
        "Robustness, beside every result that passes: the model refitted eight times, each without one rat (does p stay below 0.05 every "
        "time, and which rat does it lean on); and a linear mixed model with a random intercept for the rat (on the log for coherence, "
        "Granger and MI; on dB for power), a Wald F on k and k(n − 1) degrees of freedom, which is Kenward–Roger's for a within-rat "
        "effect in a balanced design."]))
    for mode in MODES:
        rows = significant(EM, F, mode)
        qsay = lambda q: "" if q is None else ", q < 0.001" if q < 0.001 else ", q = %.3f" % q       # noqa: E731
        txt = ["%s: %sχ²(%d) = %.2f, %s%s." % (r["label"], (r["effect"] + ", ") if r["effect"] else "", r["df"], r["chi2"], E.p_say(r["p"]),
                                              qsay(r.get("q"))) for r in rows]
        S.append(("Significant, %s" % E.MODE_SAY[mode].lower(), txt or ["Nothing passes."]))
    rob = []
    for r in (EM.get("robust") or []) + (F.get("robust") or []):
        lm = r.get("lmm") or {}
        name = r.get("name") or "Coherence, %s, %s: %s" % (E.PAIR_SAY[r["pair"]], r.get("band"), r.get("label"))
        rob.append("%s%s: leave-one-rat-out %s%s; mixed model %s." % (
            name, (" (%s)" % r["band"]) if r.get("name") and r.get("band") else "", "holds every time" if r["survives"] else "does not hold",
            "" if r["survives"] or not r.get("worst") else " (without J%d, p = %s)" % (r["worst"]["rat"], E.num(r["worst"]["p"], "%.3f")),
            ("F(%d, %d) = %.2f, %s" % (lm["df1"], lm["df2"], lm["F"], E.p_say(lm["p"]))) if lm.get("p") is not None else "not fitted"))
    S.append(("Robustness of each result", rob or ["Nothing to check."]))
    return [{"title": t, "paras": ps} for t, ps in S]


def methods_txt(M):
    lines = ["Methods, written from the final run (Jarvis tabs 11 and 12)", "=" * 58, ""]
    for s in M:
        lines += [s["title"], "-" * len(s["title"])] + [p + "\n" for p in s["paras"]]
    return "\n".join(lines)


def methods_docx(M, path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    d.add_heading("Methods, written from the final run", level=1)
    for s in M:
        d.add_heading(s["title"], level=2)
        for p in s["paras"]:
            d.add_paragraph(p)
    d.save(path)


# --------------------------------------------------------------------------
# Exports
# --------------------------------------------------------------------------
R_SCRIPT = r"""# Tab 12, the final stage: the GEEs, refitted in R (geepack: plain robust errors).
library(geepack)
d <- read.csv("final_long.csv", stringsAsFactors = TRUE)
fit <- function(sub, f, fam) {
  sub <- sub[order(sub$rat), ]; sub$rat <- factor(sub$rat)
  m <- geeglm(f, id = rat, data = sub, family = fam, corstr = "exchangeable"); print(anova(m)); invisible(m)
}
for (b in c("theta", "gamma")) for (r in c("RSC", "DHC", "ACC", "OFC", "PER")) {
  e1 <- subset(d, model == "E1" & band == b & unit == r)
  if (nrow(e1)) { cat("\nE1", b, r, "\n"); fit(e1, y ~ day * period, gaussian) }
}
for (u in levels(droplevels(subset(d, model == "GC")$unit))) { cat("\nGranger", u, "\n"); fit(subset(d, model == "GC" & unit == u), y ~ day * dir, Gamma(link = "log")) }
for (u in levels(droplevels(subset(d, model == "MI")$unit))) { cat("\nMI", u, "\n"); fit(subset(d, model == "MI" & unit == u), y ~ day, Gamma(link = "log")) }
"""


def sps_text():
    out = ["* Tab 12, the final stage: the GEEs, refitted in SPSS GENLIN (plain robust errors).",
           "GET DATA /TYPE=TXT /FILE='final_long.csv' /DELIMITERS=',' /QUALIFIER='\"' /FIRSTCASE=2",
           "  /VARIABLES=model A3 band A6 unit A24 rat F4.0 day A8 period A12 dir A8 f A4 hz F8.3 y F14.8.",
           "DATASET NAME long.", ""]
    for b in ("theta", "gamma"):
        for r in REGIONS:
            out += ["DATASET ACTIVATE long.", "TEMPORARY.", "SELECT IF (model = 'E1' AND band = '%s' AND unit = '%s')." % (b, r),
                    "GENLIN y BY day period (ORDER=ASCENDING)", "  /MODEL day period day*period INTERCEPT=YES DISTRIBUTION=NORMAL LINK=IDENTITY",
                    "  /REPEATED SUBJECT=rat WITHINSUBJECT=day*period SORT=YES CORRTYPE=EXCHANGEABLE ADJUSTCORR=YES COVB=ROBUST",
                    "  /PRINT MODELINFO FIT SUMMARY SOLUTION.", ""]
    return "\n".join(out)


def write_exports(F, PT, EM, M):
    os.makedirs(export_dir(), exist_ok=True)
    cols = ["model", "band", "unit", "rat", "day", "period", "dir", "f", "hz", "y"]
    long_rows = []
    for band, per in F["power"].items():
        for x in per:
            if not x:
                continue
            long_rows += [dict(r, model="E1", band=band, unit=x["region"]) for r in x["rows"]]
            long_rows += [dict(r, model="E2", band=band, unit=x["region"]) for r in x["rows_bins"]]
    for key, _s, _u in FOOOF_PARAMS:
        for x in F["fooof"][key]:
            long_rows += [dict(r, model="E3", band=key, unit=x["region"]) for r in x.get("rows") or []]
    for x in F["granger"]:
        long_rows += [dict(r, model="GC", band=x["band"], unit=E.PAIR_SAY[x["pair"]]) for r in x["rows"]]
    for x in F["pac"]:
        long_rows += [dict(r, model="MI", band="theta-gamma", unit=REGION_SAY[x["region"]]) for r in x["rows"]]
    files = {}
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, lineterminator="\n", extrasaction="ignore")
    w.writeheader()
    for r in long_rows:
        w.writerow({k: ("%.9g" % r[k] if k == "y" else ("%.3f" % r[k] if k == "hz" and r.get(k, "") != "" else r.get(k, ""))) for k in cols})
    files["final_long.csv"] = buf.getvalue()
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["family", "band", "unit", "test", "effect", "chi2", "df", "p", "fdr_q", "errors"])

    def row(fam, band, unit, test, x):
        if x and x.get("p") is not None:
            eff = x.get("ratio") if x.get("ratio") is not None else x.get("diff", "")
            w.writerow([fam, band, unit, test, "" if eff in ("", None) else "%.6g" % eff, "%.6g" % x["chi2"], x.get("df", 1), "%.6g" % x["p"],
                        "" if x.get("q") is None else "%.6g" % x["q"], x.get("errors", E.ROBUST)])
    for band, per in F["power"].items():
        for x in per:
            if not x:
                continue
            for t in x["e1"] or []:
                row("Power E1", band, REGION_SAY[x["region"]], t["term"], dict(t, q=x["sp"].get("q") if t["term"] == "Session × Period" else None))
            for t in x["e2"] or []:
                row("Power E2", band, REGION_SAY[x["region"]], t["term"], dict(t, q=x["spf"].get("q") if t["term"] == "Session × Period × Frequency" and x.get("spf") else None))
            for a, b, say in E.CONTRASTS:
                row("Power contrast", band, REGION_SAY[x["region"]], "Session × (%s)" % say.lower(), x["contrasts"].get(E.contrast_key(a, b)))
    for key, say, _u in FOOOF_PARAMS:
        for x in F["fooof"][key]:
            row("FOOOF", say, REGION_SAY[x["region"]], "Session × Period", x.get("sp"))
    for x in F["granger"]:
        for k, say in (("ab", "%s to %s" % (x["a"], x["b"])), ("ba", "%s to %s" % (x["b"], x["a"]))):
            row("Granger", x["band"], E.PAIR_SAY[x["pair"]], "Late vs Early, " + say, x.get(k))
        row("Granger", x["band"], E.PAIR_SAY[x["pair"]], "Session × Direction", x.get("sd"))
    for x in F["pac"]:
        row("PAC MI", "theta-gamma", REGION_SAY[x["region"]], "Late vs Early", x.get("le"))
    files["final_results.csv"] = buf.getvalue()
    # the side question is not for Joe: its own file, outside the zip
    side_buf = io.StringIO()
    w = csv.writer(side_buf, lineterminator="\n")
    w.writerow(["family", "band", "unit", "test", "effect", "chi2", "df", "p", "fdr_q", "errors"])
    for x in F["side"]:
        row("Side question (per-trial rows)", "gamma", E.PAIR_SAY[x["pair"]], "Late vs Early", x.get("le"))
    # SPSS's coefficients: the last level of each factor the reference
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["model", "band", "unit", "reference levels", "coefficient", "b", "robust_se", "chi2", "p"])
    for band, per in F["power"].items():
        for x in per:
            if not x:
                continue
            for name, rows, factors in (("E1", x["rows"], ["day", "period"]), ("E2", x["rows_bins"], ["day", "period", "f"])):
                got, _err = E._safe(lambda: spss_coefficients(E._frame(rows), factors, "normal"))
                if got:
                    refs, co = got
                    for c in co:
                        w.writerow([name, band, REGION_SAY[x["region"]], "; ".join("%s=%s" % kv for kv in refs.items()), c["term"],
                                    "%.6g" % c["b"], "%.6g" % c["se"], "%.6g" % c["chi2"], "%.6g" % c["p"]])
    files["final_coefficients_spss.csv"] = buf.getvalue()
    files["final_gee.R"] = R_SCRIPT
    files["final_gee.sps"] = sps_text()
    files["final_methods.txt"] = methods_txt(M)
    for n, text in files.items():
        with open(export_dir(n), "w", encoding="utf-8", newline="\n") as fh:
            fh.write(text)
    methods_docx(M, export_dir("final_methods.docx"))
    with open(export_dir("final_side_question_not_for_joe.csv"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(side_buf.getvalue())
    zp = export_dir("final_files.zip")
    with zipfile.ZipFile(zp + ".tmp", "w", zipfile.ZIP_DEFLATED) as z:
        for n in sorted(files) + ["final_methods.docx"]:
            z.write(export_dir(n), n)
    os.replace(zp + ".tmp", zp)
    return sorted(files) + ["final_methods.docx", "final_files.zip", "final_side_question_not_for_joe.csv"]


# --------------------------------------------------------------------------
# All of it
# --------------------------------------------------------------------------
def _clean(x, drop=("rows", "rows_bins", "G", "P")):
    if isinstance(x, dict):
        return {k: _clean(v, drop) for k, v in x.items() if k not in drop}
    if isinstance(x, list):
        return [_clean(v, drop) for v in x]
    return x


def build(progress=None, check=None):
    import shutil
    EM = E.result()
    if not EM:
        raise FinalError("Work tab 11 out first: the final stage is made from its rat-days and its hits.", 409)
    days = E.load_days()
    if not days or any("psd" not in a for a, _m in days.values()):
        raise FinalError("Tab 11's rat-days are from before the final stage: work tab 11 out again.", 409)
    shutil.rmtree(figs_dir(), ignore_errors=True)
    os.makedirs(figs_dir(), exist_ok=True)
    k = [0]

    def tick(s):
        k[0] += 1
        if check:
            check()
        if progress:
            progress(k[0], 40, s)

    def both(draw, fid):
        return draw("fdr", fid), draw("raw", fid + "_raw")
    tick("power: the spectra")
    PT = power_tables(days)
    F = {"power": {}, "fooof": {}, "granger": [], "pac": [], "side": [], "robust": []}
    figs = {"power": [], "fooof": [], "granger": [], "pac": [], "side": []}
    T = E.tables(days)
    for bid, _lo, _hi in E.BANDS:
        tick("power: the GEEs, %s" % bid)
        per = power_models(PT, bid)
        F["power"][bid] = per
        bno = 1 if bid == "theta" else 2
        for x in per:
            if not x:
                continue
            tick("power: %s, %s" % (REGION_SAY[x["region"]], bid))
            fid = "P%d_%d" % (bno, REGIONS.index(x["region"]) + 1)
            fa, fr = both(lambda m, n: fig_power_region(PT, bid, x, n, m), fid)
            cols = [("%s %s" % (d, E.PERIOD_SAY[w]), [E._c(PT["band"][w][bid][(r, d)][REGIONS.index(x["region"])]) for r in x["rats"]])
                    for d in E.DAYS for w in E.PERIODS]
            fa.update(_prism("fig_%s_prism" % fid, [{"title": "band dB by period", "rows": ["rat %d" % r for r in x["rats"]], "columns": cols,
                                                      "note": "Each rat's %s power (dB re 1 µV²/Hz) in each window." % bid}], "Figure %s" % fid))
            figs["power"].append({"id": fid, "kind": "region", "band": bid, "region": x["region"], "files": fa, "files_raw": fr,
                                  "passes": {m: E.sig(x.get("sp"), m) for m in MODES},
                                  "legend": {"title": "%s power in %s: baseline, cue 1 and cue 2 on Precon1 and Precon4." % (bid.capitalize(), REGION_SAY[x["region"]]),
                                             "body": ("(a) Band power (dB re 1 µV²/Hz), Precon1 solid, Precon4 dashed, mean ± s.e.m. of %d rats. "
                                                      "(b, c) Cue 1 − baseline and cue 2 − baseline across frequency (dB). Welch power with Dickson's "
                                                      "settings pooled over each session's trials, both sides averaged. GEE (normal, identity, "
                                                      "exchangeable, rat as the cluster): Session × Period (2 df, robust) and its planned contrasts, "
                                                      "FDR across the five regions; Session × Period × Frequency on model-based errors." % len(x["rats"]))}})
            if E.sig(x.get("sp"), "raw"):
                df = E._frame(x["rows"])
                F["robust"].append(dict(E.robustness("power", bid, None, df,
                                                     lambda d_: {"p": E._term_p(E.terms(E.gee(d_, "y ~ C(day, Sum) * C(period, Sum)", "normal")), "Session × Period")},
                                                     "y ~ C(day, Sum) * C(period, Sum)", [":"], log=False, label="Session × Period"),
                                        name="Power, %s: Session × Period" % REGION_SAY[x["region"]], test=x["sp"],
                                        passes={m: E.sig(x["sp"], m) for m in MODES}))
            for a, b, say in E.CONTRASTS:
                c = x["contrasts"].get(E.contrast_key(a, b)) or {}
                if E.sig(c, "raw"):
                    df = E._frame(x["rows"])
                    two = df[df["period"].isin([E.PERIOD_CODE[a], E.PERIOD_CODE[b]])]
                    f_ = "y ~ C(day, Treatment('Precon1')) * C(period, Treatment('%s'))" % E.PERIOD_CODE[b]
                    F["robust"].append(dict(E.robustness("power", bid, None, two,
                                                         lambda d_, f_=f_, a=a: E.interaction(d_, f_, ["Precon4", E.PERIOD_CODE[a]], "normal"),
                                                         f_, ["Precon4", E.PERIOD_CODE[a]], log=False, label="Session × (%s)" % say.lower()),
                                            name="Power, %s: Session × (%s)" % (REGION_SAY[x["region"]], say.lower()), test=c,
                                            passes={m: E.sig(c, m) for m in MODES}))
        tick("power: every region, %s" % bid)
        fid = "P%d_all" % bno
        fa, fr = both(lambda m, n: fig_power_all(PT, bid, per, n, m), fid)
        figs["power"].append({"id": fid, "kind": "all", "band": bid, "files": fa, "files_raw": fr,
                              "legend": {"title": "Within-region %s power in every region." % bid,
                                         "body": "The plan's layout: a row a region; band power by period, then cue 1 − baseline and cue 2 − baseline across frequency."}})
    tick("FOOOF: fitting")
    FT = fooof_tables(PT)
    tick("FOOOF: the GEEs")
    F["fooof"] = fooof_models(FT)
    fa, fr = both(lambda m, n: fig_fooof(F["fooof"], FT, n, m), "F1")
    figs["fooof"].append({"id": "F1", "kind": "all", "files": fa, "files_raw": fr,
                          "legend": {"title": "Oscillation against background in every region (FOOOF).",
                                     "body": ("Rows: aperiodic exponent, aperiodic offset, theta peak frequency, theta peak power above the background. "
                                              "Columns: regions. Each by period (B, C1, C2), Precon1 solid, Precon4 dashed. FOOOF 1.1, 2–55 Hz, fixed "
                                              "aperiodic mode; each side fitted, then averaged. GEE (normal) Session × Period, FDR across the regions.")}})
    tick("Granger on the hits")
    F["granger"] = granger_models(days, EM)
    for x in F["granger"]:
        if x.get("error"):
            continue
        fid = "G%d" % (len(figs["granger"]) + 1)
        fa, fr = both(lambda m, n: fig_granger(x, n, m), fid)
        figs["granger"].append({"id": fid, "kind": "pair", "band": x["band"], "pair": x["pair"], "files": fa, "files_raw": fr, "in": x["in"],
                                "passes": {m: E.sig(x.get("ab"), m) or E.sig(x.get("ba"), m) for m in MODES},
                                "legend": {"title": "Which region leads, %s, %s." % (E.PAIR_SAY[x["pair"]], x["band"]),
                                           "body": ("(a, b) Spectral Granger each way, each rat Early to Late (grey), the mean ± s.e.m. (c) Granger "
                                                    "across frequency, both ways, both days (the band shaded). From each rat-day's pooled 250 Hz "
                                                    "cross-spectral matrix (Wilson). GEE (Gamma, log link, robust): Late vs Early each way, "
                                                    "Session × Direction, FDR within the hits.")}})
        for key in ("ab", "ba"):
            if E.sig(x.get(key), "raw"):
                df = E._frame([r for r in x["rows"] if r["dir"] == ("a to b" if key == "ab" else "b to a")])
                F["robust"].append(dict(E.robustness("granger", x["band"], x["pair"], df,
                                                     lambda d_: E.ratio(E.gee(d_, "y ~ C(day, Treatment('Precon1'))"), "Precon4"),
                                                     "y ~ C(day, Treatment('Precon1'))", ["Precon4"], label="Late vs Early"),
                                        name="Granger %s → %s" % ((x["a"], x["b"]) if key == "ab" else (x["b"], x["a"])), test=x[key],
                                        passes={m: x["in"][m] and E.sig(x[key], m) for m in MODES}))
    tick("PAC in the hits' regions")
    F["pac"] = pac_models(days, EM)
    for x in F["pac"]:
        if x.get("error"):
            continue
        fid = "M%d" % (len(figs["pac"]) + 1)
        fa, fr = both(lambda m, n: fig_pac(x, n, m), fid)
        figs["pac"].append({"id": fid, "kind": "region", "region": x["region"], "files": fa, "files_raw": fr, "in": x["in"],
                            "passes": {m: E.sig(x.get("le"), m) for m in MODES},
                            "legend": {"title": "Theta–gamma coupling within %s, Early against Late." % REGION_SAY[x["region"]],
                                       "body": ("(a) Tort's modulation index, each rat (grey) and the mean ± s.e.m. (b) Gamma amplitude across theta "
                                                "phase (each bin's share), Early solid, Late dashed. Pooled over every trial's cue pair, the same wire; "
                                                "a region the mean of its sides. GEE (Gamma, log link, robust) Late vs Early.")}})
        if E.sig(x.get("le"), "raw"):
            F["robust"].append(dict(E.robustness("pac", None, None, E._frame(x["rows"]),
                                                 lambda d_: E.ratio(E.gee(d_, "y ~ C(day, Treatment('Precon1'))"), "Precon4"),
                                                 "y ~ C(day, Treatment('Precon1'))", ["Precon4"], label="Late vs Early"),
                                    name="Theta–gamma MI, %s" % REGION_SAY[x["region"]], test=x["le"],
                                    passes={m: x["in"][m] and E.sig(x["le"], m) for m in MODES}))
    tick("the side question")
    F["side"] = side_question(days, T)
    fa, fr = both(lambda m, n: fig_side(F["side"], EM, n, m), "S1")
    figs["side"].append({"id": "S1", "kind": "all", "files": fa, "files_raw": fr, "joe": False,
                         "legend": {"title": "Side question (not for Joe): gamma coherence with every trial its own row.",
                                    "body": "Each trial's own coherence (the mean of its valid wire pairs) a GEE row, the rat the cluster; beside it, tab 11's pooled answer."}})
    tick("the methods")
    M = methods(EM, F, _sampling_rate())
    out = {"schema": SCHEMA, "at": time.strftime("%Y-%m-%dT%H:%M:%S"), "from": EM.get("at"), "rats": EM["rats"],
           "regions": [{"id": r, "say": REGION_SAY[r], "color": REGION_COLORS[r]} for r in REGIONS],
           "hits": {m: [{"pair": p, "band": b} for p, b in hits(EM, m)] for m in MODES},
           "hit_regions": {m: hit_regions(EM, m) for m in MODES},
           "power": {b: [_clean(x) for x in per] for b, per in F["power"].items()},
           "fooof": {k: [_clean(x) for x in v] for k, v in F["fooof"].items()}, "fooof_params": [list(p) for p in FOOOF_PARAMS],
           "granger": [_clean(x) for x in F["granger"]], "pac": [_clean(x) for x in F["pac"]], "side": [_clean(x) for x in F["side"]],
           "robust": F["robust"], "figures": figs, "methods": M, "twenty": TWENTY, "modes": {m: E.MODE_SAY[m] for m in MODES}}
    out["exports"] = write_exports(F, PT, EM, M)
    joe._write(_root("final.json"), out)
    return out


def result():
    p = _root("final.json")
    if not os.path.isfile(p):
        return None
    with open(p, encoding="utf-8") as fh:
        return json.load(fh)


def work(worker):
    """Tab 11 again first when its rat-days are from before the final
    stage's measures, then everything here."""
    days = E.load_days()
    if not days or any("psd" not in a for a, _m in days.values()):
        E.work(worker)
    out = build(progress=lambda i, n, s: worker.note(phase="final", i=i, of=n, item=s), check=worker.check)
    return {"figures": sum(len(v) for v in out["figures"].values())}


def work_safe(worker):
    from . import monolith as MO
    try:
        return work(worker)
    except (FinalError, E.EvenMoreError, joe.JoeError) as exc:
        raise MO.MonolithError(str(exc), exc.status)


if __name__ == "__main__":
    from . import monolith as MO
    MO.configure(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "GUI_logs"))
    o = build(progress=lambda i, n, s: print("%d/%d %s" % (i, n, s), flush=True))
    print(sum(len(v) for v in o["figures"].values()), "figures")
